# %%


from __future__ import annotations

import math
import warnings
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd

try:
    import matplotlib
    import matplotlib.pyplot as plt
except Exception:
    matplotlib = None
    plt = None


# %%
# PART 1 — CONFIG
########################################################################################
########################################################################################

FIG_DIR: str | None = "figures"      # None이면 저장 안 함
FIG_DPI = 130
SHOW_FIGURES = True

FAST_MODE = True

# ── 운전 조건 ─────────────────────────────────────────────────────────────


T_OPER = 353.15          # 운전 온도 [K] (80 C)
P_AN = 1.0               # 양극 압력 [bar]
P_CAT = 30.0             # 음극 압력 [bar]
J0_CA = 0.2              # 음극 교환전류밀도 [A/cm2]
ALPHA_CA = 0.5           # 음극 전달계수
MEMB_T_UM = 180.0        # 막 두께 [um] (~Nafion 117)
MEMB_LAMBDA = 22.0       # 막 함수율 lambda
SIGMA_SCALE = 1.0        # Springer sigma 배율
J_LIM = 6.1              # 물질전달 한계 전류밀도 [A/cm2]
BOP_EFFICIENCY = 4.2     # BOP 소비 [kWh/kg-H2]

STACK_SIZING_MODE = "nameplate_kW"   # "nameplate_kW" | "pv_fraction" | "cells"
STACK_NAMEPLATE_KW = 1000.0         # PEM 명판 [kW]
STACK_PV_FRACTION = 0.70
STACK_N_CELLS = 200.0
A_CELL_CM2 = 1000.0       # 셀 활성면적 [cm2]
N_CELL_PER_STACK = 150.0   # 스택당 셀 수
INTEGER_STACKS = False    # 정수 스택 수로 올림하지 않음

J_MAX = 2.0              # 정격 전류밀도 [A/cm2]
J_X_mAcm2 = 2.0          # 크로스오버 등가 전류밀도 [mA/cm2]
J_MIN_SAFETY = 2.0       # 안전계수. (1.0 -> j_min=50*j_x, 2.0 -> 100*j_x)
J_MIN = 50.0 * J_MIN_SAFETY * (J_X_mAcm2 / 1000.0)

J_STAR_MODE = "optimize"                       # "fixed" | "optimize"
J_STAR_FIXED = 1.80                         # 공통 고정 전류밀도 [A/cm2]
J_STAR_FIXED_CASE2: float | None = None     # None이면 J_STAR_FIXED 사용
J_STAR_FIXED_CASE3: float | None = None     # None이면 J_STAR_FIXED 사용

RE_SOURCE = "region"

# 스크립트와 같은 폴더에 "지역별 발전소 데이터" 폴더를 두면 됩니다.
# 다른 위치에 보관한다면 RE_DATA_ROOT만 실제 절대경로로 변경하세요.
PROJECT_DIR = Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
RE_DATA_ROOT = str(PROJECT_DIR / "지역별 발전소 데이터" / "1MW 정규화"
                   / "지역별 태양광, 풍력")

RE_REGION_NAME = "세종"                   # 파일명에서 지역 자동 탐색: available_regions()로 확인
RE_KIND = "태양광"                        # "태양광" | "풍력"
# 태양광: 강원 경기 경남 경북 광주 대구 대전 부산 서울 세종 울산 인천 전남 전북 제주 충남 충북
# 풍력:   강원 경기 경남 경북 부산 울산 인천 전남 전북 제주 충남

RE_CAPACITY_MW = 3                       # 발전량 정규화 규모 [MW]
RE_VALUE_COL = "1MW_정규화_발전량(kWh)"
RE_DATETIME_COL = "datetime"

RE_FILE_PATH: str | None = None
RE_TIME_COL: str | None = None
RE_POWER_COL: str | None = None

RE_YEAR_RANGE_BY_KIND = {"태양광": (2018, 2025), "풍력": (2018, 2022)}
RE_START_YEAR: int | None = None          # 시작연도 (포함); None -> 발전원별 기본값
RE_END_YEAR: int | None = None            # 종료연도 (포함); None -> 발전원별 기본값
RE_MISSING_POLICY = "error"              # 결측은 기존처럼 오류. 실제 숫자 0과 구분합니다.
RE_MIN_ANNUAL_CF = 0.01                  # 기존의 연간 이용률 1% 점검 기준 유지
RE_LOW_CF_POLICY = "warn"                # "warn": 경고 후 원값 사용 | "error": 기존처럼 중단
# 부산 풍력 2021~2022년은 업로드 파일의 발전량 자체가 모두 0입니다.

RE_FILE_VALUE_UNIT = "kW"                # RE_SOURCE="file" 전용: "kW" | "kWh"
SAVE_RESULT_CSV = True

# 최종 요약 Excel 저장: FIG_DIR / SAVE_RESULT_CSV / 그래프 생성 여부와 독립적입니다.
SAVE_FINAL_SUMMARY_XLSX = True
FINAL_SUMMARY_DIR: str | Path = "results_excel"  

RE_REGION = "(unset)"
RE_SOURCE_URL = ""
ALLOCATION_N = 1.00                      # 배정 비율 n (0~1)
DT_HOURS = 1.0

ELEC_PRICE = 0.089       # 전기 단가 [$/kWh]
WATER_PRICE = 0.0018     #  [$/kg]
LIFETIME_YR = 20         # 프로젝트 기간 [yr]
INTEREST_RATE = 0.08     # 할인율
PEM_FIXED_OM_FRAC = 0.05   # PEM (stack + BOP) CAPEX의 5%/yr
ESS_FIXED_OM_FRAC = 0.015  # ESS CAPEX의 1.5%/yr
VARIABLE_OM = 0.024      #  [$/kg-H2]
STACK_LIFETIME_H = 40000.0   # 스택 수명 상한 [h]
REPL_STACK_FRAC = 1.0       # 교체비 / 스택 CAPEX


# ── 스택 열화 · 교체  ──────────────────────────────────────────────

DEG_J_REF = 1.80             # [A/cm2] 앵커 기준 운전점
STACK_EOL_RETENTION = 0.90   # EOL 성능 유지율 (BOL 대비)
DEG_DJ_REF = 0.40            # [A/cm2/h]
DEG_SEVERITY_CAP = 2.5       # severity 상한
DEG_K_START = 2.0            # OFF -> ON 기동 열화 배율
DEG_J_EPS_STOP = 0.15        # [A/cm2] ON -> OFF 열화 계산용 proxy

DEG_J_EXTRAPOLATE = "linear"  # "linear" | "clamp" | "error"; j > 3 A/cm2 처리

REPLACEMENT_DOWNTIME_H = 168.0   # 교체 1회당 정지 [h]
DOWNTIME_DUTY_WEIGHTED = True
ELEC_PAID_DURING_DOWNTIME = True

MANUFACTURING_RATE = 100.0   # 생산규모 [MW/yr] (BOP 단가 보간용)
MARKUP = 0.5                 # 제조원가 대비 마크업
OVHD_FRAC = 0.5              # 스택 오버헤드 비율

ELEC_PENALTY_MULT = 2.0      # 위약금 배율 (무차원)

ETA_RTE = 0.9                # 왕복효율. 확정 범위 0.85~0.90
C_E_USD_KWH = 310           # ESS 에너지 단가 [$/kWh]
C_P_USD_KW = 310           # ESS 전력 단가 [$/kW]
ANNUAL_BATT_REPL_USD = 0.0    # 연간 배터리 교체 충당금 [USD/yr]
CHARGE_RATING_POLICY = "p95"  # "raw" | "p95" | "p99" | ...
CHARGE_RATING_PCTL = 95.0     # "raw" 계열이 아닐 때만 사용


def _use_charge_pctl() -> bool:
    return str(CHARGE_RATING_POLICY).strip().lower() not in ("raw", "peak", "max", "none")

_ESS_AUTO_EXTEND_Q = True
SURPLUS_ETA_CHG = 1.0             # None -> sqrt(ETA_RTE) (편도 충전효율)
SURPLUS_DAY_HOURS = 24.0          # 일별 잉여전력 적분 주기 [h]

CASE1 = dict(
    rating_sweep_factors=(0.4, 0.6, 0.8, 1.0, 1.3, 1.6, 2.0),   # 명판 kW 배수
)
CASE2 = dict(
    soc_max=0.95,
    soc_floor=0.05,                # Case 2는 정지선 = floor -> usable 0.90
    control_interval_h=1.0,        # 기동 판단 간격 [h]
    charge_while_running=True,    # 집합 ESS의 동시 입·출력 포트 근사
                                 # False: 충·방전 시간 분리
    q_grid=(2, 5, 10, 15, 20, 25, 30, 35, 40, 50, 60, 70, 80, 90, 100),
    q_grid_fast=(2, 5, 10, 20, 30, 40, 60, 80, 100),
    n_j_grid=40, n_j_grid_fast=16,
)
CASE3 = dict(
    soc_max=0.95,             # SOC 상한
    soc_floor=0.05,           # 안전 floor
    soc_stop=0.05,            # 실질 정지선
    soc_restart=0.15,         # 재기동선
    restart_hours_of_Pstar=None,
    direct_restart=True,
    q_grid=(2, 5, 10, 15, 20, 25, 30, 35, 40, 50, 60, 70, 80, 90, 100),
    q_grid_fast=(2, 5, 10, 20, 30, 40, 60, 80, 100),
    n_j_grid=40, n_j_grid_fast=20,
)

INVERSION_GRID_N = 4000
WARM_START = True                  # 전체 분석 시작에 한 번만 적용; 워밍업 실적 제외
WARM_START_DAYS = 30.0              # 마지막 관측 30일을 초기상태 산정에만 사용
DEG_EVENT_REFERENCE_H = 1.0        # 이벤트 손상 환산 기준 [h], 실제 기동시간 아님


def _n_j_grid(case: str) -> int:
    d = {"case2": CASE2, "case3": CASE3}[case]
    return d["n_j_grid_fast"] if FAST_MODE else d["n_j_grid"]

def _q_grid(case: str):
    d = {"case2": CASE2, "case3": CASE3}[case]
    k = "q_grid_fast" if FAST_MODE else "q_grid"
    return d[k]

def j_is_fixed() -> bool:
    return str(J_STAR_MODE).lower() == "fixed"

def fixed_j(case: str) -> float:
    per = {"case2": J_STAR_FIXED_CASE2, "case3": J_STAR_FIXED_CASE3}[case]
    raw = per if per is not None else J_STAR_FIXED
    try:
        j = float(raw)
    except (TypeError, ValueError):
        raise ValueError(f"{case}: 전류밀도는 숫자여야 합니다 -> {raw!r}") from None
    if not np.isfinite(j) or j <= 0:
        raise ValueError(f"{case}: 전류밀도는 유한한 양수여야 합니다 -> {j}")
    return float(np.clip(j, J_MIN, J_MAX))

_OUTPUT_RUN_TAG = ""
_FIGURE_MANIFEST = []


def _ensure_fig_dir():
    if FIG_DIR is None:
        return None
    d = _resolve_path(FIG_DIR)
    if _OUTPUT_RUN_TAG:
        d = d / _OUTPUT_RUN_TAG
    d.mkdir(parents=True, exist_ok=True)
    return str(d)


def daily_storage_need(P_in, P_star, P_rc, dt=1.0, eta_c=None,
                       threshold=None, day_hours=None) -> np.ndarray:
    P = np.asarray(P_in, dtype=float)
    if eta_c is None:
        eta_c = (math.sqrt(ETA_RTE) if SURPLUS_ETA_CHG is None
                 else float(SURPLUS_ETA_CHG))

    threshold = float(P_star) if threshold is None else float(threshold)
    day_hours = SURPLUS_DAY_HOURS if day_hours is None else day_hours

    spd = int(round(day_hours / dt))                    # steps per day
    if spd < 1:
        raise ValueError(f"day_hours({day_hours}) / dt({dt}) 가 1스텝 미만입니다.")
    n_days = len(P) // spd
    if n_days < 1:
        raise ValueError(f"프로파일 길이 {len(P)} 가 하루({spd}스텝)보다 짧습니다.")
    P = P[:n_days * spd] 

    sur = np.maximum(P - threshold, 0.0)                # 초과분만 저장 대상
    sur = np.minimum(sur, max(float(P_rc), 0.0))        # 충전정격 클리핑
    return eta_c * sur.reshape(n_days, spd).sum(axis=1) * dt

def size_ess_from_surplus(P_in, P_star, P_rc, q, dt=1.0, case="case3",
                          soc_max=None, soc_stop=None, eta_RTE=None,
                          return_need=False):
    d = {"case2": CASE2, "case3": CASE3}[case]
    soc_max = d["soc_max"] if soc_max is None else soc_max
    if soc_stop is None:
        soc_stop = d.get("soc_stop", d.get("soc_floor", 0.0))
    usable = float(soc_max) - float(soc_stop)
    if usable <= 0:
        raise ValueError(f"{case}: SOC 사용 창이 0 이하입니다 "
                         f"(max {soc_max} - stop {soc_stop}).")

    eta = ETA_RTE if eta_RTE is None else eta_RTE
    thr = float(P_star) if case == "case3" else float(P_star) / eta

    e_draw = daily_storage_need(P_in, P_star, P_rc, dt=dt, threshold=thr)
    E_rated = float(np.percentile(e_draw, q)) / usable
    return (E_rated, e_draw) if return_need else E_rated


# %%
# PART 2 — 상수 & 원가 데이터
########################################################################################
########################################################################################

R_CONST = 8.314462618        # 기체상수 [J/mol/K]
F_CONST = 96485.33212        # 패러데이 상수 [C/mol]
MW_H2 = 2.01588e-3           # 수소 분자량 [kg/mol]
MW_H2O = 18.01528e-3         # 물 분자량 [kg/mol]
N_ELECTRON_H2 = 2.0          # 수소 1몰당 전자수
EUR_TO_USD = 1.17
WATER_STOICH = MW_H2O / MW_H2    # 화학량론 물 소요 [kg/kg-H2]
HHV_H2_KWH_KG = 39.39            # 수소 고위발열량 [kWh/kg] (141.8 MJ/kg)

METAL_PRICE = {
    "Pt": 33487.0 / 1000.0,
    "Au": 49.000 * EUR_TO_USD,
    "Ir": 122906.0/1000.0,
}

MATERIAL_PRICE = {
    "membrane_180um": 766.0 * EUR_TO_USD,   # Nafion 180 um [USD/m2]
    "membrane_80um":  371.0 * EUR_TO_USD,   # Nafion 80 um [USD/m2]
    "carbon_cloth":   168.0 * EUR_TO_USD,   # 음극 bulk PTL [USD/m2]
    "PPS40GF":        15.4 * EUR_TO_USD,    # 프레임/가스켓 [USD/kg]
    "ti_powder":      0.400 * EUR_TO_USD,   # 양극 bulk PTL [USD/g]
    "ss316":          0.0013 * EUR_TO_USD,  # BPP [USD/g]
    "ptfe":           8.0 * EUR_TO_USD,     # 프레임/가스켓 [USD/kg]
    "a356_al":        82.0 * EUR_TO_USD,    # 엔드플레이트 [USD/m2]
}

DENSITY = {"Ti": 4.50, "SS316": 8.00, "Au": 19.32, "Pt": 21.45}

CA_LOAD_mgcm2 = 0.2      # 음극 Pt 로딩
CA_CAT = "Pt"
PTL_AN_MAT = "Ti"        # 양극 PTL 소재
PTL_AN_T_mm = 1.5
PTL_AN_COAT = "Pt"       # 양극 PTL 코팅
PTL_AN_COAT_T_um = 0.1
PTL_CA_CC = 1.0          # 음극 carbon cloth 계수
BPP_T_mm = 1.0
BPP_COAT = "Au"
BPP_COAT_T_um = 0.1
FRAME_GASKET = 1.0
EP_AREA = 1.25           # 엔드플레이트 면적 계수

BOP_USD_KW = {10.0: 575.0, 100.0: 520.0, 1000.0: 377.0}


def _log_interpolate(x_value: float, table: dict) -> float:
    xs = np.array(sorted(table), dtype=float)
    ys = np.array([table[x] for x in xs], dtype=float)
    if x_value <= xs[0]:
        return float(ys[0])
    if x_value >= xs[-1]:
        lx = np.log10(xs[-2:])
        slope = (ys[-1] - ys[-2]) / (lx[1] - lx[0])
        return float(ys[-1] + slope * (np.log10(x_value) - lx[1]))
    return float(np.interp(np.log10(x_value), np.log10(xs), ys))

def _mem_cost_per_m2(thickness_um: float) -> float:
    x1, y1 = 180.0, MATERIAL_PRICE["membrane_180um"]
    x2, y2 = 80.0, MATERIAL_PRICE["membrane_80um"]
    price = y1 + (y2 - y1) * (thickness_um - x1) / (x2 - x1)
    return max(price, 0.30 * y2)

def _mass_material_g(area_cm2: float, thickness_um: float, density_gcm3: float) -> float:
    return area_cm2 * (thickness_um * 1.0e-4) * density_gcm3


# %%
# PART 3 — 촉매 : 고정 Tafel 파라미터
########################################################################################
########################################################################################

@dataclass(frozen=True)
class IrO2Params:
    slope_mV_dec: float = 50.0    # Tafel 기울기 [mV/dec]
    j0_Acm2: float = 1.0e-6      # 교환전류밀도 [A/cm2]
    an_load_mgcm2: float = 1.0   # Ir 질량 기준 로딩 [mg/cm2]

    @property
    def b_V_dec(self) -> float:
        return self.slope_mV_dec / 1000.0

    def eta(self, j_Acm2):
        j = np.clip(np.atleast_1d(np.asarray(j_Acm2, dtype=float)), 1e-30, None)
        return self.b_V_dec * np.log10(j / self.j0_Acm2)

IRO2 = IrO2Params()


# %%
# PART 4 — PEM 셀 전압 모델
########################################################################################
########################################################################################

def membrane_conductivity(T_K: float | None = None,                 # Springer model
                          lam: float | None = None,
                          scale: float | None = None) -> float:
    T_K = T_OPER if T_K is None else T_K
    lam = MEMB_LAMBDA if lam is None else lam
    scale = SIGMA_SCALE if scale is None else scale
    sigma_30 = 0.005139 * lam - 0.00326
    sigma = sigma_30 * math.exp(1268.0 * (1.0 / 303.0 - 1.0 / T_K))
    return max(scale * sigma, 1.0e-4)

def cell_voltage_parts(j_Acm2) -> dict:
    j = np.atleast_1d(np.asarray(j_Acm2, dtype=float))
    j = np.clip(j, 1.0e-30, None)
    if np.any(j >= J_LIM):
        raise ValueError(f"전류밀도가 j_lim({J_LIM})에 도달했습니다 — ln(1-j/j_lim) 발산")

    T = T_OPER
    p_sat = 10.0 ** (35.4462 - 3343.93 / T - 10.9 * np.log10(T) + 4.1645e-3 * T)
    p_h2 = max(P_CAT - p_sat, 1.0e-12)
    p_o2 = max(P_AN - p_sat, 1.0e-12)
    U0 = 1.5184 - 1.5421e-3 * T + 9.523e-5 * T * np.log(T) + 9.84e-8 * T ** 2
    e_rev = U0 + (R_CONST * T / (2.0 * F_CONST)) * np.log(p_h2 * np.sqrt(p_o2))

    eta_act_an = IRO2.eta(j)

    eta_act_ca = (R_CONST * T / (N_ELECTRON_H2 * ALPHA_CA * F_CONST)) \
        * np.arcsinh(j / (2.0 * J0_CA))

    memb_asr = (MEMB_T_UM * 1.0e-4) / membrane_conductivity()
    eta_ohm = j * memb_asr

    eta_mt = -(R_CONST * T / (N_ELECTRON_H2 * F_CONST)) * np.log(1.0 - j / J_LIM)

    return {"j": j, "e_rev": np.full_like(j, e_rev), "eta_act_an": eta_act_an,
            "eta_act_ca": eta_act_ca, "eta_ohm": eta_ohm, "eta_mt": eta_mt,
            "memb_asr": memb_asr,
            "V": e_rev + eta_act_an + eta_act_ca + eta_ohm + eta_mt}


# %%
# PART 5 — 물리 커널 래퍼 + P <-> j 역산
########################################################################################
########################################################################################

def faraday_effi(j):
    # Faraday 효율 = 1.0
    return np.ones_like(np.asarray(j, dtype=float))

def V(j):
    j = np.atleast_1d(np.asarray(j, dtype=float))
    out = np.zeros_like(j)
    on = j > 0.0
    if np.any(on):
        out[on] = cell_voltage_parts(j[on])["V"]
    return out

def SEC_stack(j):
    return V(j) * 2.0 * F_CONST / MW_H2 / 3.6e6

def SEC_system(j):
    return SEC_stack(j) + BOP_EFFICIENCY

def h2_area(j):
    j = np.atleast_1d(np.asarray(j, dtype=float))
    return faraday_effi(j) * j / (2.0 * F_CONST) * MW_H2 * 3600.0

def p_sys(j):
    j = np.atleast_1d(np.asarray(j, dtype=float))
    return j * V(j) / 1000.0 + BOP_EFFICIENCY * h2_area(j)

def P_of_j(j, A_tot):
    return A_tot * p_sys(j)

def split_power(j, A_tot):
    j = float(np.atleast_1d(j)[0])
    P_stack = A_tot * j * float(V(j)[0]) / 1000.0
    P_bop = A_tot * BOP_EFFICIENCY * float(h2_area(j)[0])
    return P_stack, P_bop

@dataclass
class OperatingWindow:
    j_min: float
    j_max: float
    P_min: float
    P_max: float
    A_tot: float

    @property
    def current_turndown(self) -> float:
        return self.j_min / self.j_max

    @property
    def power_turndown(self) -> float:
        return self.P_min / self.P_max

    def describe(self) -> str:
        return (f"  j window : {self.j_min:.4f} ~ {self.j_max:.4f} A/cm2\n"
                f"  P window : {self.P_min:,.1f} ~ {self.P_max:,.1f} kW\n"
                f"  turndown : current {self.current_turndown * 100:.2f} %  vs  "
                f"power {self.power_turndown * 100:.2f} %")

def operating_window(A_tot: float, j_min: float | None = None,
                     j_max: float | None = None) -> OperatingWindow:
    j_min = J_MIN if j_min is None else j_min
    j_max = J_MAX if j_max is None else j_max
    if j_max >= J_LIM:
        raise ValueError(f"j_max({j_max}) must be < j_lim({J_LIM})")
    return OperatingWindow(j_min, j_max,
                           float(P_of_j(j_min, A_tot)[0]),
                           float(P_of_j(j_max, A_tot)[0]), A_tot)

def _kernel_signature() -> tuple:
    return (T_OPER, MEMB_T_UM, MEMB_LAMBDA, SIGMA_SCALE, J_LIM, BOP_EFFICIENCY,
            IRO2.slope_mV_dec, IRO2.j0_Acm2, J0_CA, ALPHA_CA, P_AN, P_CAT)

class PowerToJ:

    def __init__(self, A_tot: float, j_lo: float, j_hi: float, n: int | None = None):
        self.A_tot, self.j_lo, self.j_hi = float(A_tot), float(j_lo), float(j_hi)
        self.n = int(n or INVERSION_GRID_N)
        self._jg = np.linspace(self.j_lo, self.j_hi, self.n)
        self._pg = self.A_tot * p_sys(self._jg)
        if not np.all(np.diff(self._pg) > 0):
            raise RuntimeError("p_sys(j)가 단조증가하지 않아 역산할 수 없습니다")

    def __call__(self, P):
        return np.interp(np.asarray(P, dtype=float), self._pg, self._jg)

_INVERTER_CACHE: dict = {}

def get_inverter(A_tot: float, j_lo: float, j_hi: float,       
                 n: int | None = None) -> PowerToJ:
    key = (round(A_tot, 6), round(j_lo, 8), round(j_hi, 8),
           int(n or INVERSION_GRID_N), _kernel_signature())
    inv = _INVERTER_CACHE.get(key)
    if inv is None:
        inv = PowerToJ(A_tot, j_lo, j_hi, n)
        _INVERTER_CACHE[key] = inv
    return inv


# %%
# PART 6 — 스택 원가 + 고정 스택 사이징
########################################################################################
########################################################################################

def stack_material_costs(N_total_cells: float, N_stacks: float,
                         A_cell_cm2: float | None = None) -> dict:
    A_cell_cm2 = A_CELL_CM2 if A_cell_cm2 is None else A_cell_cm2
    area_cm2 = N_total_cells * A_cell_cm2
    area_m2 = area_cm2 / 1.0e4

    membrane = _mem_cost_per_m2(MEMB_T_UM) * area_m2

    an_cat_g = IRO2.an_load_mgcm2 * area_cm2 / 1000.0
    an_cat = an_cat_g * METAL_PRICE["Ir"] * 1.3

    ca_cat_g = CA_LOAD_mgcm2 * area_cm2 / 1000.0
    ca_cat = ca_cat_g * METAL_PRICE[CA_CAT] * 1.3

    ptl_an_bulk = _mass_material_g(area_cm2, PTL_AN_T_mm * 1000.0,
                                   DENSITY[PTL_AN_MAT]) * MATERIAL_PRICE["ti_powder"]
    ptl_an_coat = _mass_material_g(area_cm2, PTL_AN_COAT_T_um,
                                   DENSITY[PTL_AN_COAT]) * METAL_PRICE[PTL_AN_COAT]
    ptl_ca = area_m2 * MATERIAL_PRICE["carbon_cloth"] * PTL_CA_CC

    bpp_bulk = _mass_material_g(area_cm2, BPP_T_mm * 1000.0,
                                DENSITY["SS316"]) * MATERIAL_PRICE["ss316"]
    bpp_coat = _mass_material_g(area_cm2, BPP_COAT_T_um,
                                DENSITY[BPP_COAT]) * METAL_PRICE[BPP_COAT]

    gasket_kg = area_m2 * 2.0 * FRAME_GASKET
    frame_gasket = (0.7 * gasket_kg * MATERIAL_PRICE["PPS40GF"]
                    + 0.3 * gasket_kg * MATERIAL_PRICE["ptfe"])

    endplate = N_stacks * 2.0 * (A_cell_cm2 / 1.0e4) * EP_AREA * MATERIAL_PRICE["a356_al"]

    total = (membrane + an_cat + ca_cat + ptl_an_bulk + ptl_an_coat + ptl_ca
             + bpp_bulk + bpp_coat + frame_gasket + endplate)
    return {"membrane": membrane, "anode_catalyst": an_cat, "cathode_catalyst": ca_cat,
            "ptl_anode_bulk": ptl_an_bulk, "ptl_anode_coating": ptl_an_coat,
            "ptl_cathode": ptl_ca, "bpp_bulk": bpp_bulk, "bpp_coating": bpp_coat,
            "frame_gasket": frame_gasket, "endplate": endplate,
            "total_material_cost": total, "total_active_area_m2": area_m2}

def stack_total_cost(N_total_cells: float, N_stacks: float,
                     A_cell_cm2: float | None = None) -> dict:
    mat = stack_material_costs(N_total_cells, N_stacks, A_cell_cm2)
    direct = mat["total_material_cost"] * (8.0 + 4.0 + 1.0) / 8.0
    return {**mat, "direct_stack_cost": direct,
            "total_stack_cost": direct / (1.0 - OVHD_FRAC)}

def bop_total_capex(P_plant_kW: float) -> dict:
    usd_kW = _log_interpolate(MANUFACTURING_RATE, BOP_USD_KW)
    manufactured = usd_kW * P_plant_kW
    return {"bop_usd_kW": usd_kW, "bop_manufactured": manufactured,
            "bop_total": manufactured * (1.0 + MARKUP)}


@dataclass
class FixedStack:
    A_cell: float
    N_cells: float
    N_stacks: float
    A_tot: float
    j_rated: float
    P_nameplate: float
    P_stack_rated: float
    P_bop_rated: float
    capex_stack: float
    capex_bop: float
    cost_breakdown: dict = field(default_factory=dict)

    @property
    def capex_fix(self) -> float:
        return self.capex_stack + self.capex_bop

    def report(self) -> None:
        print("\n[PART 6] fixed stack sizing")
        print(f"  cells={self.N_cells:,.1f} ({self.N_stacks:.2f} stacks), "
              f"A_tot={self.A_tot:,.0f} cm2, j_rated={self.j_rated:.3f} A/cm2")
        print(f"  nameplate={self.P_nameplate:,.1f} kW "
              f"(stack={self.P_stack_rated:,.1f}, BOP={self.P_bop_rated:,.1f} kW)")
        print(f"  CAPEX stack={self.capex_stack:,.0f}, BOP={self.capex_bop:,.0f}, "
              f"fixed={self.capex_fix:,.0f} USD")

def size_fixed_stack(P_target_kW: float, j_rated: float | None = None) -> FixedStack:
    j_rated = J_MAX if j_rated is None else float(j_rated)
    p_cell_plant = float(p_sys(j_rated)[0]) * A_CELL_CM2

    N_cells = P_target_kW / p_cell_plant

    N_stacks = N_cells / N_CELL_PER_STACK
    if INTEGER_STACKS:
        N_stacks = math.ceil(N_stacks)
        N_cells = N_stacks * N_CELL_PER_STACK

    A_tot = N_cells * A_CELL_CM2
    P_stack_r, P_bop_r = split_power(j_rated, A_tot)
    P_nameplate = P_stack_r + P_bop_r

    cost = stack_total_cost(N_cells, N_stacks, A_CELL_CM2)
    return FixedStack(A_cell=A_CELL_CM2, N_cells=N_cells, N_stacks=N_stacks, A_tot=A_tot,
                      j_rated=j_rated, P_nameplate=P_nameplate,
                      P_stack_rated=P_stack_r, P_bop_rated=P_bop_r,
                      capex_stack=float(cost["total_stack_cost"]),
                      capex_bop=float(bop_total_capex(P_nameplate)["bop_total"]),
                      cost_breakdown=cost)

def build_fixed_stack(profile=None, verbose: bool = True) -> FixedStack:
    if STACK_SIZING_MODE == "nameplate_kW":
        P_target, basis = STACK_NAMEPLATE_KW, f"명판 직접 지정 {STACK_NAMEPLATE_KW:,.0f} kW"
    elif STACK_SIZING_MODE == "pv_fraction":
        if profile is None:
            raise ValueError("'pv_fraction' 모드는 profile이 필요합니다.")
        P_target = STACK_PV_FRACTION * profile.peak
        basis = f"배정 PV 피크 {profile.peak:,.0f} kW x {STACK_PV_FRACTION:.2f}"
    elif STACK_SIZING_MODE == "cells":
        P_target = STACK_N_CELLS * A_CELL_CM2 * float(p_sys(J_MAX)[0])
        basis = f"셀 수 직접 지정 {STACK_N_CELLS:,.0f} cells"
    else:
        raise ValueError(f"Unknown STACK_SIZING_MODE: {STACK_SIZING_MODE}")

    st = size_fixed_stack(P_target)
    if verbose:
        print(f"\n[sizing] mode={STACK_SIZING_MODE}, target={P_target:,.1f} kW ({basis})")
        st.report()
    return st

def rippl_capacity(P_in: np.ndarray, P_star: float, eta_RTE: float,
                   dt: float = 1.0) -> float:
    C = np.cumsum((np.asarray(P_in, dtype=float) - P_star / eta_RTE) * dt)
    return float(C.max() - C.min())


# %%
# PART 7 — 재생에너지 프로파일 로더
########################################################################################
########################################################################################

HOURS_PER_YEAR = 8760 
_TIME_HINTS = ("time", "timestamp", "date", "datetime", "일시", "시간", "날짜")
_POWER_HINTS = ("kw", "power", "gen", "output", "발전", "출력", "전력", "kwh")


@dataclass
class REProfile:
    P_in: np.ndarray
    P_gen: np.ndarray
    timestamps: pd.DatetimeIndex
    dt: float = 1.0
    allocation_n: float = 1.0
    region: str = ""
    year: int = 0                  # 단일연도 호출 호환; 표시에는 period_label 사용
    source: str = ""
    url: str = ""
    meta: dict = field(default_factory=dict)

    def __post_init__(self):
        self.P_in = np.asarray(self.P_in, dtype=float)
        self.P_gen = np.asarray(self.P_gen, dtype=float)
        self.timestamps = pd.DatetimeIndex(self.timestamps)
        if (self.P_in.ndim != 1 or not len(self.P_in) or
                len(self.P_gen) != len(self.P_in) or len(self.timestamps) != len(self.P_in)):
            raise ValueError("P_in, P_gen, timestamps의 길이가 같고 비어 있지 않아야 합니다")
        if (not np.isfinite(self.dt) or self.dt <= 0 or
                np.any(~np.isfinite(self.P_in)) or np.any(~np.isfinite(self.P_gen)) or
                np.any(self.P_in < 0) or np.any(self.P_gen < 0)):
            raise ValueError("전력은 유한한 비음수, dt는 양수여야 합니다")
        if self.timestamps.hasnans or not self.timestamps.is_monotonic_increasing or self.timestamps.has_duplicates:
            raise ValueError("시간축에 결측/중복/역순이 있습니다")
        if len(self.timestamps) > 1:
            gaps = np.asarray((self.timestamps[1:] - self.timestamps[:-1]).total_seconds()) / 3600
            if not np.allclose(gaps, self.dt, rtol=0, atol=1e-7):
                raise ValueError("시간축은 dt 간격의 연속 시계열이어야 합니다")
        self.year = int(self.timestamps[0].year)

    @property
    def years(self) -> tuple[int, ...]:
        return tuple(int(y) for y in self.timestamps.year.unique())

    @property
    def n_years(self) -> int:
        return len(self.years)

    @property
    def period_label(self) -> str:
        return str(self.years[0]) if self.n_years == 1 else f"{self.years[0]}–{self.years[-1]}"

    @property
    def n_hours(self) -> int:
        return len(self.P_in)

    @property
    def E_paid(self) -> float:
        return float(self.P_in.sum() * self.dt)

    @property
    def peak(self) -> float:
        return float(self.P_in.max())

    @property
    def mean(self) -> float:
        return float(self.P_in.mean())

    @property
    def capacity_factor(self) -> float:
        cap_kw = float(self.meta.get("re_capacity_MW", 0.0) or 0.0) * 1000.0
        basis = cap_kw if cap_kw > 0 else float(self.P_gen.max())
        return float(self.P_gen.mean()) / basis if basis > 0 else 0.0

    @property
    def month(self) -> np.ndarray:
        return self.timestamps.month.to_numpy()

    def pctl(self, q: float) -> float:
        return float(np.percentile(self.P_in, q))

    def daily_energy(self) -> np.ndarray:
        return pd.Series(self.P_in * self.dt, index=self.timestamps).resample("D").sum().to_numpy()

    def representative_days(self) -> dict:
        s = pd.Series(self.P_in * self.dt, index=self.timestamps).resample("D").sum()
        pool = s[s > s.max() * 0.05]
        low = pool.idxmin() if len(pool) else s.idxmin()
        dates = self.timestamps.normalize()
        return {"summer_peak": np.flatnonzero(dates == s.idxmax()),
                "winter_low": np.flatnonzero(dates == low)}

    def report(self) -> None:
        print("\n[PART 7] renewable profile")
        print(f"  {self.region} / {self.period_label} ({self.n_years} yr) | {self.source}")
        print(f"  steps={self.n_hours:,}, dt={self.dt:g} h, allocation={self.allocation_n:.0%}, "
              f"mean annual E_paid={self.E_paid / self.n_years / 1e3:,.1f} MWh")
        print(f"  peak/mean={self.peak:,.1f}/{self.mean:,.1f} kW, CF={self.capacity_factor:.2%}")


def _normalize_re_name(value) -> str:
    import unicodedata
    return unicodedata.normalize("NFC", str(value)).strip()


def _analysis_years(start_year=None, end_year=None, *, kind=None) -> tuple[int, ...]:
    selected_kind = _normalize_re_name(RE_KIND if kind is None else kind)
    if selected_kind not in RE_YEAR_RANGE_BY_KIND:
        raise ValueError(f"지원하지 않는 발전원: {selected_kind!r}. '태양광' 또는 '풍력'을 사용하세요.")
    default_start, default_end = RE_YEAR_RANGE_BY_KIND[selected_kind]
    start = RE_START_YEAR if start_year is None else start_year
    end = RE_END_YEAR if end_year is None else end_year
    start = int(default_start if start is None else start)
    end = int(default_end if end is None else end)
    if not 1900 <= start <= end <= 2200:
        raise ValueError("1900 <= 분석 시작연도 <= 분석 종료연도 <= 2200이 필요합니다")
    return tuple(range(start, end + 1))


def _resolve_path(path: str | Path) -> Path:
    p = Path(path).expanduser()
    if p.is_absolute():
        return p
    base = Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
    candidate = base / p
    return candidate if candidate.exists() or not p.exists() else p.resolve()


def _pick_column(df, hints, exclude=()):
    return next((c for c in df.columns if c not in exclude and
                 any(h in str(c).strip().lower() for h in hints)), None)


def _read_input_table(path: Path) -> pd.DataFrame:
    if path.suffix.lower() in (".xlsx", ".xls"):
        return pd.read_excel(path)
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return pd.read_csv(path, encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"CSV 인코딩을 읽을 수 없습니다: {path}")


def _parse_timestamps(values) -> pd.DatetimeIndex:
    if pd.api.types.is_numeric_dtype(values):
        ts = pd.DatetimeIndex(pd.to_datetime(values, unit="D", origin="1899-12-30", errors="coerce"))
    else:
        ts = pd.DatetimeIndex(pd.to_datetime(values, errors="coerce"))
    if ts.tz is not None:
        ts = ts.tz_convert("Asia/Seoul").tz_localize(None)
    rounded = ts.round("s")
    delta = np.asarray(abs(ts - rounded).total_seconds())
    if np.all(np.isnan(delta) | (delta <= 0.001)):
        ts = rounded
    return ts


def _validated_series(frames, years, label, dt, missing_policy=None):
    policy = str(RE_MISSING_POLICY if missing_policy is None else missing_policy).lower()
    if policy not in ("error", "zero"):
        raise ValueError("RE_MISSING_POLICY는 'error' 또는 'zero'여야 합니다")
    frame = pd.concat(frames, ignore_index=True)
    if frame["timestamp"].isna().any():
        raise ValueError(f"{label}: 해석할 수 없는 datetime이 있습니다")
    frame = frame.loc[frame.timestamp.dt.year.isin(years)].copy()
    missing_years = sorted(set(years) - set(frame.timestamp.dt.year))
    if missing_years:
        raise ValueError(f"{label}: 요청 연도 자료가 없습니다: {missing_years}. "
                         "자료를 추가하거나 시작/종료 연도를 확인하세요.")
    if len(frame) == 0:
        raise ValueError(f"{label}: 선택한 기간의 데이터가 없습니다")
    bad = ~np.isfinite(frame["value"].to_numpy(dtype=float))
    frame.loc[bad, "value"] = np.nan
    dup = frame.timestamp.duplicated(keep=False)
    removed = 0
    if dup.any():
        groups = frame.loc[dup].groupby("timestamp")["value"]
        for t, vals in groups:
            a = vals.to_numpy(dtype=float)
            if not (np.isnan(a).all() or (np.isfinite(a).all() and np.allclose(a, a[0], rtol=1e-10, atol=1e-8))):
                raise ValueError(f"{label}: 중복 시각의 발전량이 다릅니다: {t}")
        removed = int(frame.timestamp.duplicated().sum())
        frame = frame.drop_duplicates("timestamp")
    step = pd.Timedelta(hours=float(dt))
    if dt <= 0 or not np.isclose(24 / dt, round(24 / dt)):
        raise ValueError("DT_HOURS는 양수이고 24시간을 나누어야 합니다")
    expected = pd.date_range(f"{years[0]}-01-01", f"{years[-1] + 1}-01-01", freq=step, inclusive="left")
    series = frame.set_index("timestamp")["value"].sort_index()
    off_grid = series.index.difference(expected)
    if len(off_grid):
        raise ValueError(f"{label}: 설정된 시간 간격과 맞지 않는 시각 {off_grid[0]}")
    missing = expected.difference(series.index)
    series = series.reindex(expected)
    invalid = series.isna()
    if invalid.any() and policy == "error":
        example = ", ".join(str(t) for t in series.index[invalid][:3])
        raise ValueError(f"{label}: 결측 {int(invalid.sum()):,}개 (누락 시각 {len(missing):,}개). "
                         f"예: {example}. 자동 0 보정하지 않습니다.")
    if (series.dropna() < 0).any():
        raise ValueError(f"{label}: 음수 발전량이 있습니다")
    quality = []
    for y in years:
        mask = expected.year == y
        quality.append(dict(year=y, expected_steps=int(mask.sum()),
                            missing_timestamps=int((missing.year == y).sum()),
                            invalid_values=int(invalid[mask].sum()),
                            filled_zero=int(invalid[mask].sum()) if policy == "zero" else 0))
    if invalid.any():
        warnings.warn(f"{label}: 명시한 zero 정책으로 {int(invalid.sum()):,}개를 0 보정합니다", RuntimeWarning)
        series = series.fillna(0.0)
    return series, {"quality": quality, "duplicate_rows_removed": removed, "missing_policy": policy}


def _region_file_catalog(kind=None, data_root=None) -> dict[str, list[Path]]:
    """지역별 1MW 정규화 파일만 탐색. 원본/발전소별 파일과 혼합하지 않음."""
    selected_kind = _normalize_re_name(RE_KIND if kind is None else kind)
    if selected_kind not in RE_YEAR_RANGE_BY_KIND:
        raise ValueError(f"지원하지 않는 발전원: {selected_kind!r}")
    root = _resolve_path(RE_DATA_ROOT if data_root is None else data_root)
    if not root.is_dir():
        raise FileNotFoundError(f"데이터 폴더 없음: {root}. RE_DATA_ROOT를 확인하세요.")
    catalog: dict[str, list[Path]] = {}
    for path in root.rglob("*"):
        if (not path.is_file() or path.suffix.lower() not in (".xlsx", ".xls", ".csv")
                or path.name.startswith("~$")):
            continue
        name = _normalize_re_name(path.name)
        parts = name.split("_")
        if (len(parts) < 3 or parts[1] != selected_kind or
                "1MW" not in name or "정규화" not in name or "원본" in name):
            continue
        catalog.setdefault(parts[0], []).append(path)
    return {region: sorted(paths, key=lambda p: _normalize_re_name(str(p)))
            for region, paths in sorted(catalog.items())}


def available_regions(kind=None, data_root=None) -> tuple[str, ...]:
    return tuple(_region_file_catalog(kind, data_root))


def find_region_files(region=None, kind=None, data_root=None) -> list[Path]:
    region = _normalize_re_name(RE_REGION_NAME if region is None else region)
    kind = _normalize_re_name(RE_KIND if kind is None else kind)
    catalog = _region_file_catalog(kind, data_root)
    if region not in catalog:
        root = _resolve_path(RE_DATA_ROOT if data_root is None else data_root)
        raise FileNotFoundError(
            f"{region} {kind} 1MW 정규화 파일 없음. 검색 폴더: {root}. "
            f"사용 가능한 {kind} 지역: {', '.join(catalog) or '(없음)'}")
    return catalog[region]


def find_region_file(region=None, kind=None, data_root=None) -> str:
    paths = find_region_files(region, kind, data_root)
    if len(paths) != 1:
        raise ValueError("여러 입력 파일이 있습니다. find_region_files()를 사용하세요.")
    return str(paths[0])


def load_profile_from_region(region=None, kind=None, year=None, re_mw=None,
                             data_root=None, verbose=True, *, start_year=None, end_year=None):
    region = _normalize_re_name(RE_REGION_NAME if region is None else region)
    kind = _normalize_re_name(RE_KIND if kind is None else kind)
    if year is not None:
        if start_year is not None or end_year is not None:
            raise ValueError("year와 start_year/end_year를 동시에 지정할 수 없습니다")
        start_year = end_year = int(year)
    years = _analysis_years(start_year, end_year, kind=kind)
    re_mw = float(RE_CAPACITY_MW if re_mw is None else re_mw)
    if re_mw <= 0 or not np.isfinite(re_mw) or not 0 <= ALLOCATION_N <= 1:
        raise ValueError("재생 용량은 양수, ALLOCATION_N은 0~1이어야 합니다")
    frames, paths = [], find_region_files(region, kind, data_root)
    used_paths = []
    for path in paths:
        df = _read_input_table(path)
        for col in (RE_DATETIME_COL, RE_VALUE_COL):
            if col not in df.columns:
                raise KeyError(f"{path.name}: '{col}' 컬럼이 없습니다")
        ts = _parse_timestamps(df[RE_DATETIME_COL])
        frames.append(pd.DataFrame({"timestamp": ts, "value": pd.to_numeric(df[RE_VALUE_COL], errors="coerce")}))
        if np.any(np.isin(ts.year, years)):
            used_paths.append(str(path))
    s, meta = _validated_series(frames, years, f"{region} {kind}", DT_HOURS)
    p_unit_kw = s.to_numpy(dtype=float) / DT_HOURS
    cf_year = pd.Series(p_unit_kw, index=s.index).groupby(s.index.year).mean() / 1000.0
    policy = str(RE_LOW_CF_POLICY).strip().lower()
    threshold = float(RE_MIN_ANNUAL_CF)
    if policy not in ("warn", "error"):
        raise ValueError("RE_LOW_CF_POLICY는 'warn' 또는 'error'여야 합니다")
    if not np.isfinite(threshold) or not 0 <= threshold <= 1:
        raise ValueError("RE_MIN_ANNUAL_CF는 0~1 사이의 유한한 숫자여야 합니다")
    low_years = [int(y) for y in cf_year.index[cf_year < threshold]]
    zero_years = [int(y) for y in cf_year.index[cf_year == 0]]
    if not np.any(p_unit_kw > 0):
        raise ValueError(f"{region} {kind}: 선택한 전 기간의 발전량이 0입니다. 원자료를 확인하세요.")
    if low_years:
        message = (f"{region} {kind}: 연간 이용률 {threshold:.1%} 미만 연도 {low_years}. "
                   f"연간 발전량이 모두 0인 연도: {zero_years}. 원자료 확인이 필요합니다.")
        if policy == "error":
            raise ValueError(message)
        warnings.warn(message + " 원값을 변경하거나 연도를 제외하지 않고 계산합니다.",
                      RuntimeWarning, stacklevel=2)
    for row in meta["quality"]:
        y = row["year"]
        row.update(capacity_factor=float(cf_year.loc[y]), low_capacity_factor=y in low_years,
                   zero_generation_year=y in zero_years)
    meta.update(re_capacity_MW=re_mw, value_col=RE_VALUE_COL, unit_CF=float(p_unit_kw.mean() / 1000),
                paths=used_paths, kind=kind, requested_years=years, full_years=True,
                annual_capacity_factor={int(y): float(cf) for y, cf in cf_year.items()},
                low_cf_years=low_years, zero_generation_years=zero_years,
                low_cf_policy=policy, min_annual_cf=threshold)
    P_gen = p_unit_kw * re_mw
    prof = REProfile(P_in=P_gen * ALLOCATION_N, P_gen=P_gen, timestamps=s.index, dt=DT_HOURS,
                     allocation_n=ALLOCATION_N, region=f"{region} {kind}",
                     source=f"{len(used_paths)} file(s) × {re_mw:g} MW", url=RE_SOURCE_URL, meta=meta)
    if verbose:
        print(f"[RE loader] {region} {kind} {prof.period_label}: {len(P_gen):,} steps, "
              f"{prof.n_years} yr, CF={prof.capacity_factor:.2%}")
    return prof


def load_profile_from_file(path, time_col=None, power_col=None, allocation_n=None,
                           *, start_year=None, end_year=None) -> REProfile:
    p = _resolve_path(path)
    if not p.is_file():
        raise FileNotFoundError(p)
    df = _read_input_table(p)
    tcol = time_col or _pick_column(df, _TIME_HINTS)
    pcol = power_col or _pick_column(df, _POWER_HINTS, exclude=(tcol,))
    if tcol is None or pcol is None:
        raise ValueError("다년 분석 파일에는 datetime과 발전량 컬럼이 필요합니다")
    years = _analysis_years(start_year, end_year)
    frame = pd.DataFrame({"timestamp": _parse_timestamps(df[tcol]),
                          "value": pd.to_numeric(df[pcol], errors="coerce")})
    s, meta = _validated_series([frame], years, p.name, DT_HOURS)
    unit = str(RE_FILE_VALUE_UNIT).lower()
    if unit not in ("kw", "kwh"):
        raise ValueError("RE_FILE_VALUE_UNIT은 'kW' 또는 'kWh'여야 합니다")
    power = s.to_numpy() / (DT_HOURS if unit == "kwh" else 1.0)
    n = float(ALLOCATION_N if allocation_n is None else allocation_n)
    if not 0 <= n <= 1:
        raise ValueError("allocation_n은 0~1이어야 합니다")
    meta.update(paths=[str(p)], kind=RE_KIND, full_years=True, time_col=tcol, power_col=pcol)
    return REProfile(power * n, power, s.index, DT_HOURS, n, RE_REGION,
                     source=f"file:{p.name}", url=RE_SOURCE_URL, meta=meta)


def _profile_signature() -> tuple:
    return (RE_SOURCE, RE_REGION_NAME, RE_KIND, _analysis_years(), RE_CAPACITY_MW,
            ALLOCATION_N, RE_FILE_PATH, RE_DATA_ROOT, DT_HOURS, RE_MISSING_POLICY,
            RE_VALUE_COL, RE_DATETIME_COL, RE_TIME_COL, RE_POWER_COL, RE_FILE_VALUE_UNIT,
            RE_LOW_CF_POLICY, RE_MIN_ANNUAL_CF)


_PROFILE_CACHE: REProfile | None = None
_PROFILE_SIG: tuple | None = None


def get_profile(force_reload=False) -> REProfile:
    global _PROFILE_CACHE, _PROFILE_SIG
    sig = _profile_signature()
    if _PROFILE_CACHE is not None and not force_reload and _PROFILE_SIG == sig:
        return _PROFILE_CACHE
    if RE_SOURCE == "region":
        prof = load_profile_from_region()
    elif RE_SOURCE == "file" and RE_FILE_PATH:
        prof = load_profile_from_file(RE_FILE_PATH, RE_TIME_COL, RE_POWER_COL)
    else:
        raise ValueError("RE_SOURCE/RE_FILE_PATH 설정을 확인하세요")
    _PROFILE_CACHE, _PROFILE_SIG = prof, sig
    return prof


# %%
# PART 7.5 — PEM 열화 커널
########################################################################################
########################################################################################

_SU_CONST_J = np.array([1.0, 2.0, 3.0])            # [A/cm2]
_SU_CONST_UV = np.array([22.7, 26.1, 50.0])        # [uV/h] constant current (C1/C2/C3)
_SU_SOLAR_UV = np.array([39.7, 52.4, 87.7])        # [uV/h] solar PV mode
_SU_J_VALID = 3.0                                  # Su 커널 유효 상한 [A/cm2]


def sanitize_operating_j(j):
    """운전 창 적용: 0 < j < J_MIN은 OFF, 상한은 J_LIM."""
    a = np.asarray(j, dtype=float).copy()
    if np.any(~np.isfinite(a)):
        raise ValueError("전류밀도 프로파일에 NaN/inf 가 있습니다")
    ceil = float(J_LIM)
    tol = max(1e-6, 1e-9 * ceil)
    over = a.max(initial=0.0) - ceil
    if over > tol:
        raise ValueError(f"전류밀도 {a.max():.6f} 가 물질전달 한계 J_LIM={ceil:.2f} "
                         f"A/cm2 를 {over:.3e} 만큼 초과합니다")
    a[a < 0.0] = 0.0
    a[(a > 0.0) & (a < J_MIN)] = 0.0
    a = np.minimum(a, ceil)
    return float(a) if a.ndim == 0 else a


_SU_TOP_SLOPE = float((_SU_CONST_UV[2] - _SU_CONST_UV[1])
                      / (_SU_CONST_J[2] - _SU_CONST_J[1]))       # 23.9 uV/h per A/cm2


def su_degradation(j: float) -> float:
    """Su 정전류 열화율 [uV/h]. j > 3은 지정 외삽 규칙 적용."""
    j = float(j)
    if j <= 0.0:
        return 0.0
    if j < 1.0:
        return float(_SU_CONST_UV[0]) * j                 # 원점 -> C1 선형
    if j <= _SU_J_VALID + 1e-9:
        return float(np.interp(j, _SU_CONST_J, _SU_CONST_UV))

    mode = str(DEG_J_EXTRAPOLATE).strip().lower()
    if mode == "error":
        raise ValueError(f"j={j:.3f} A/cm2 는 Su 커널 유효범위(0~3 A/cm2) 밖입니다 "
                         f"(DEG_J_EXTRAPOLATE='error')")
    if mode == "clamp":
        return float(_SU_CONST_UV[2])
    return float(_SU_CONST_UV[2]) + _SU_TOP_SLOPE * (j - _SU_J_VALID)


# --- 벡터화 버전 (8760 h 적산용. 위 스칼라 함수와 수치적으로 동일) ---

def _su_degradation_vec(j: np.ndarray) -> np.ndarray:
    j = np.asarray(j, dtype=float)
    out = np.where(j < 1.0, _SU_CONST_UV[0] * j,
                   np.interp(np.clip(j, _SU_CONST_J[0], _SU_CONST_J[-1]),
                             _SU_CONST_J, _SU_CONST_UV))
    j_top = float(j.max(initial=0.0))
    if j_top > _SU_J_VALID + 1e-9:                        # j > 3 외삽
        mode = str(DEG_J_EXTRAPOLATE).strip().lower()
        if mode == "error":
            raise ValueError(f"j={j_top:.3f} A/cm2 는 Su 커널 유효범위(0~3) 밖입니다 "
                             f"(DEG_J_EXTRAPOLATE='error')")
        if mode != "clamp":
            extra = np.maximum(j - _SU_J_VALID, 0.0) * _SU_TOP_SLOPE
            out = out + extra
    return np.where(j <= 0.0, 0.0, out)

def _su_solar_factor_vec(j: np.ndarray) -> np.ndarray:
    return np.interp(np.clip(np.asarray(j, dtype=float),
                             _SU_CONST_J[0], _SU_CONST_J[-1]),
                     _SU_CONST_J, _SU_SOLAR_UV / _SU_CONST_UV)

def _dynamic_factor_vec(j: np.ndarray, delta_j: np.ndarray) -> np.ndarray:
    j = np.asarray(j, dtype=float)
    sev = np.clip(np.abs(np.asarray(delta_j, dtype=float)) / max(DEG_DJ_REF, 1e-9),
                  0.0, DEG_SEVERITY_CAP)
    return np.where(j <= 0.0, 1.0, 1.0 + (_su_solar_factor_vec(j) - 1.0) * sev)

def stack_eol_delta_v(j_ref: float | None = None) -> float:
    """EOL 임계전압 [V] = V_BOL × (1 / retention - 1)."""
    j_ref = DEG_J_REF if j_ref is None else float(j_ref)
    j_ref = float(np.clip(j_ref, J_MIN, J_LIM))
    v0 = float(np.atleast_1d(V(j_ref))[0])
    return v0 * (1.0 / max(float(STACK_EOL_RETENTION), 1e-9) - 1.0)

def deg_anchor_rate_uV_h(j_ref: float | None = None) -> float:
    """앵커 열화율 [uV/h] = EOL 임계전압 / 스택 수명."""
    return stack_eol_delta_v(j_ref) / max(float(STACK_LIFETIME_H), 1.0) * 1.0e6

def deg_anchor_scale(j_ref: float | None = None) -> float:
    """Su 열화율을 기준 전류밀도·스택 수명에 맞춰 환산합니다."""
    j_ref = DEG_J_REF if j_ref is None else float(j_ref)
    j_ref = float(np.clip(j_ref, J_MIN, J_LIM))
    return deg_anchor_rate_uV_h(j_ref) / max(su_degradation(j_ref), 1e-12)

def degradation_summary(j_profile, dt=1.0, initial_j: float = 0.0, analysis_years=None) -> dict:
    """정상 손상은 운전시간, 기동·정지·변화 손상은 이벤트별 적산."""
    j = np.atleast_1d(sanitize_operating_j(j_profile)).astype(float)
    if j.ndim != 1 or len(j) == 0:
        raise ValueError("전류밀도 프로파일은 비어 있지 않은 1차원 배열이어야 합니다")
    h = np.broadcast_to(np.asarray(dt, dtype=float), j.shape).copy()
    if np.any(~np.isfinite(h)) or np.any(h < 0) or h.sum() <= 0:
        raise ValueError("구간 길이는 유한한 비음수이고 총 기간은 양수여야 합니다")
    if DEG_K_START < 1 or DEG_EVENT_REFERENCE_H <= 0:
        raise ValueError("DEG_K_START >= 1, DEG_EVENT_REFERENCE_H > 0이 필요합니다")
    scale = deg_anchor_scale()
    prev = np.r_[float(sanitize_operating_j(initial_j)), j[:-1]]
    on, was_on = j > 0, prev > 0
    is_start, is_stop, is_on = on & ~was_on, ~on & was_on, on & was_on
    dj = np.abs(j - prev)
    base = _su_degradation_vec(j) * scale
    ref_h = float(DEG_EVENT_REFERENCE_H)
    steady = base * h
    start = np.where(is_start, base * (DEG_K_START - 1.0) * ref_h, 0.0)
    ramp = np.where(is_on, base * (_dynamic_factor_vec(j, dj) - 1.0) * ref_h, 0.0)
    jp = float(min(DEG_J_EPS_STOP, J_MAX))
    stop = np.where(is_stop, su_degradation(jp) * scale *
                    _dynamic_factor_vec(np.full(len(j), jp), prev - jp) * ref_h, 0.0)
    delta_uV = steady + start + ramp + stop
    dv = delta_uV * 1e-6
    cum = np.cumsum(dv)
    total_v = float(cum[-1])
    period_h = float(h.sum())
    profile_years = float(analysis_years) if analysis_years is not None else period_h / HOURS_PER_YEAR
    op_h_period = float(h[on].sum())
    op_h_year = op_h_period / profile_years
    annual_v = total_v / profile_years
    effective = total_v / op_h_period * 1e6 if op_h_period > 0 else 0.0
    j_mean = (float(np.dot(j, h) / op_h_period) if op_h_period > 0
              else float(np.clip(DEG_J_REF, J_MIN, J_LIM)))
    eol = stack_eol_delta_v(j_mean)
    life_voltage = eol / (effective * 1e-6) if effective > 0 else float("inf")
    life_h = min(float(STACK_LIFETIME_H), life_voltage)
    binding = (f"voltage (+{(1.0 / STACK_EOL_RETENTION - 1.0) * 100:.1f} %)"
               if life_voltage < STACK_LIFETIME_H else f"hours ({STACK_LIFETIME_H:,.0f} h)")
    contrib = {"steady": float(steady.sum()) * 1e-6,
               "ramp": float(ramp.sum()) * 1e-6,
               "start": float(start.sum()) * 1e-6,
               "stop": float(stop.sum()) * 1e-6}
    state_h = {"OFF": float(h[~on & ~was_on].sum()),
               "START": float(h[is_start].sum()), "STOP": float(h[is_stop].sum()),
               "ON_ON": float(h[is_on].sum())}
    rate = np.divide(delta_uV, h, out=np.zeros_like(h), where=h > 0)
    return {
        "rate_eff_uV_h": effective, "rate_raw_uV_h": effective / max(scale, 1e-12),
        "anchor_scale": scale, "annual_degradation_V": annual_v,
        "annual_degradation_mV": annual_v * 1e3, "annual_op_hours": op_h_year,
        "j_mean_operating": j_mean,
        "delta_j_mean": (float(np.average(dj[is_on], weights=h[is_on]))
                         if h[is_on].sum() > 0 else 0.0),
        "eol_delta_V": eol, "stack_life_hours": life_h,
        "stack_life_hours_voltage": life_voltage,
        "stack_life_years": life_h / op_h_year if op_h_year > 0 else float("inf"),
        "life_binding": binding, "starts": int(is_start.sum()), "stops": int(is_stop.sum()),
        "state_hours": state_h, "contrib_V": contrib,
        "contrib_pct": {k: v / total_v * 100 if total_v > 0 else 0.0 for k, v in contrib.items()},
        "dr_uV_h": rate, "delta_uV": delta_uV, "cumulative_V_profile": cum,
        "duration_h": h, "profile_hours": period_h, "profile_years": profile_years,
        "reference_event_h": ref_h, "start_flags": is_start, "stop_flags": is_stop,
        "contrib_uV_series": {"steady": steady, "ramp": ramp, "start": start, "stop": stop},
        "annual_starts": int(is_start.sum()) / profile_years,
        "annual_stops": int(is_stop.sum()) / profile_years,
        "annual_contrib_V": {k: v / profile_years for k, v in contrib.items()},
    }

def _deg_config_fingerprint() -> tuple:
    return (DEG_J_REF, DEG_DJ_REF, DEG_SEVERITY_CAP, DEG_K_START,
            DEG_J_EPS_STOP, DEG_J_EXTRAPOLATE, STACK_LIFETIME_H, STACK_EOL_RETENTION,
            J_MIN, J_MAX, J_LIM, DEG_EVENT_REFERENCE_H, _kernel_signature())

def ensure_degradation(res) -> dict:
    fp = _deg_config_fingerprint()
    if res.meta.get("degradation_fp") != fp:
        j = res.meta.get("segment_j", res.j)
        durations = res.meta.get("segment_duration_h", res.dt)
        res.meta["degradation"] = degradation_summary(
            j, dt=durations, initial_j=res.meta.get("initial_j", 0.0),
            analysis_years=res.meta.get("n_years"))
        res.meta["degradation_fp"] = fp
    return res.meta["degradation"]

def print_degradation_summary(results: dict) -> pd.DataFrame:
    rows = []
    for name, res in results.items():
        d = ensure_degradation(res)
        c = d["contrib_pct"]
        rows.append({
            "case": name,
            "평균 j [A/cm2]": d["j_mean_operating"],
            "평균 |dj|": d["delta_j_mean"],
            "유효 열화율 [uV/h]": d["rate_eff_uV_h"],
            "연 열화 [mV/yr]": d["annual_degradation_mV"],
            "스택 수명 [h]": d["stack_life_hours"],
            "스택 수명 [yr]": d["stack_life_years"],
            "수명 결정": d["life_binding"],
            "기동/yr": d["annual_starts"],
            "정상 [%]": c["steady"], "램프 [%]": c["ramp"],
            "기동 [%]": c["start"], "정지 [%]": c["stop"],
        })
    df = pd.DataFrame(rows)
    print(f"\n[PEM 열화 · 스택 수명]  Su et al. 2024 커널 "
          f"(앵커={STACK_LIFETIME_H:,.0f} h @ {STACK_EOL_RETENTION * 100:.0f} %)")
    print(df.to_string(index=False, float_format=lambda x: f"{x:,.2f}"))
    return df


# %%
# PART 8 — 결과 스키마 · 손실 원장 · 지표
########################################################################################
########################################################################################

@dataclass
class DispatchResult:
    case: str
    dt: float = 1.0
    P_in: np.ndarray = field(default_factory=lambda: np.zeros(0))
    j: np.ndarray = field(default_factory=lambda: np.zeros(0))
    P_used: np.ndarray = field(default_factory=lambda: np.zeros(0))
    H2_series: np.ndarray = field(default_factory=lambda: np.zeros(0))
    soc: np.ndarray | None = None
    running: np.ndarray | None = None
    charge: np.ndarray | None = None
    discharge: np.ndarray | None = None
    curtail_series: np.ndarray | None = None
    E_paid: float = 0.0
    E_curtail: float = 0.0
    E_idle: float = 0.0
    E_rte: float = 0.0
    E_bop: float = 0.0
    E_stack: float = 0.0
    E_direct: float = 0.0
    E_via_ess: float = 0.0
    H2_total: float = 0.0
    op_hours: float = 0.0
    hours_curtail: float = 0.0
    hours_idle: float = 0.0
    hours_night: float = 0.0
    restarts: int = 0
    eq_cycles: float = 0.0
    E_start: float = 0.0           
    E_end: float = 0.0
    on_hours: np.ndarray | None = None 
    E_rated: float = 0.0
    P_rated_chg: float = 0.0
    P_rated_dis: float = 0.0
    j_star: float | None = None
    P_star: float | None = None
    meta: dict = field(default_factory=dict)
    timestamps: pd.DatetimeIndex | None = None

    @property
    def E_used(self) -> float:
        return self.E_bop + self.E_stack

    @property
    def E_unabsorbed(self) -> float:
        return self.E_curtail + self.E_idle

    @property
    def ledger(self) -> dict:
        return {"curtail": self.E_curtail, "idle": self.E_idle, "rte": self.E_rte,
                "bop": self.E_bop, "stack": self.E_stack,
                "storage_delta": self.E_end - self.E_start}

    def SEC_eff(self) -> float:
        return self.E_paid / self.H2_total if self.H2_total > 0 else float("nan")

    def SEC_used(self) -> float:
        return self.E_used / self.H2_total if self.H2_total > 0 else float("nan")

    def SEC_stack_avg(self) -> float:
        return self.E_stack / self.H2_total if self.H2_total > 0 else float("nan")

    def util(self, h2_area_at: float, A_tot: float, n_hours=None) -> float:
        n_hours = n_hours or len(self.P_in) or HOURS_PER_YEAR
        d = h2_area_at * A_tot * n_hours * self.dt
        return self.H2_total / d if d > 0 else float("nan")

    def ledger_table(self) -> pd.DataFrame:
        labels = {"curtail": "낭비 · curtailment", "idle": "낭비 · 하한미만 정지",
                  "rte": "낭비 · ESS 왕복손실", "bop": "실사용 · BOP(오버헤드)",
                  "stack": "실사용 · stack(생산)",
                  "storage_delta": "ESS 저장량 변화 (종료 - 시작)"}
        tot = self.E_paid
        rows = [{"bucket": labels[k], "MWh": v / 1e3,
                 "share_%": (v / tot * 100) if tot > 0 else np.nan}
                for k, v in self.ledger.items()]
        rows.append({"bucket": "합계 (= 배정 입력량)", "MWh": tot / 1e3, "share_%": 100.0})
        return pd.DataFrame(rows)

    def report(self, h2_area_at_rated=None, h2_area_at_jstar=None, A_tot=None) -> None:
        n = len(self.P_in) or HOURS_PER_YEAR
        print(f"\n[{self.case}] mean annual H2={_annual_mean(self, self.H2_total) / 1000:,.2f} t | "
              f"SEC_eff={self.SEC_eff():.2f} kWh/kg | op={_annual_mean(self, self.op_hours):,.0f} h/yr")
        if self.E_paid > 0:
            print(f"  ledger share: curtail={self.E_curtail / self.E_paid * 100:.2f} %, "
                  f"idle={self.E_idle / self.E_paid * 100:.2f} %, "
                  f"rte={self.E_rte / self.E_paid * 100:.2f} %, "
                  f"used={self.E_used / self.E_paid * 100:.2f} %")
        if A_tot and h2_area_at_rated:
            print(f"  util vs rated={self.util(h2_area_at_rated, A_tot, n) * 100:.2f} %")
        if A_tot and h2_area_at_jstar:
            print(f"  util vs j*={self.util(h2_area_at_jstar, A_tot, n) * 100:.2f} %")
        if self.E_rated > 0:
            print(f"  ESS: E={self.E_rated:,.0f} kWh, "
                  f"Pchg/Pdis={self.P_rated_chg:,.0f}/{self.P_rated_dis:,.0f} kW, "
                  f"restarts/yr={_annual_mean(self, self.restarts):,.1f}, cycles/yr={_annual_mean(self, self.eq_cycles):,.1f}, "
                  f"SOC_end={self.E_end / self.E_rated * 100:.1f} %")
        if self.E_via_ess > 0 or self.E_direct > 0:
            tot = self.E_direct + self.E_via_ess
            print(f"  period path: direct={self.E_direct / 1e3:,.1f} MWh "
                  f"({self.E_direct / tot * 100:.1f} %), "
                  f"via_ESS={self.E_via_ess / 1e3:,.1f} MWh "
                  f"({self.E_via_ess / tot * 100:.1f} %)")


def compare_table(results: dict, lcohs: dict | None = None) -> pd.DataFrame:
    rows = []
    for name, r in results.items():
        row = {"case": name, "mean_annual_H2 [t]": _annual_mean(r, r.H2_total) / 1000.0,
               "SEC_eff [kWh/kg]": r.SEC_eff(), "op_hours [h/yr]": _annual_mean(r, r.op_hours),
               "curtail [%]": r.E_curtail / r.E_paid * 100 if r.E_paid else np.nan,
               "idle [%]": r.E_idle / r.E_paid * 100 if r.E_paid else np.nan,
               "RTE loss [%]": r.E_rte / r.E_paid * 100 if r.E_paid else np.nan,
               "E_rated [kWh]": r.E_rated, "eq_cycles": _annual_mean(r, r.eq_cycles),
               "restarts": _annual_mean(r, r.restarts)}
        a = annual_result_table(r)
        row["annual_H2_std [t]"] = float(a.H2_kg.std(ddof=1)) / 1e3 if len(a) > 1 else 0.0
        row["worst_annual_H2 [t]"] = float(a.H2_kg.min()) / 1e3
        row["source_years"] = len(a)
        row["worst_source_year"] = int(a.H2_kg.idxmin())
        row["j_star [A/cm2]"] = r.j_star if r.j_star is not None else np.nan
        row["P_star [kW]"] = r.P_star if r.P_star is not None else np.nan
        row["P_chg [kW]"] = r.P_rated_chg
        row["P_dis [kW]"] = r.P_rated_dis
        deg = ensure_degradation(r)
        row["열화율 [uV/h]"] = deg["rate_eff_uV_h"]
        row["스택 수명 [h]"] = deg["stack_life_hours"]
        row["수명 결정"] = deg["life_binding"]
        if lcohs and name in lcohs:
            row["LCOH [$/kg]"] = lcohs[name]["LCOH"]
            row["ESS share [%]"] = lcohs[name].get("ess_share_pct", np.nan)
            row["교체 [회]"] = lcohs[name].get("replacement_count", np.nan)
            row["생애평균 유지율 [%]"] = lcohs[name].get("retention_mean_pct", np.nan)
        rows.append(row)
    return pd.DataFrame(rows)


def _bind_result(res, timestamps=None):
    if timestamps is not None:
        ts = pd.DatetimeIndex(timestamps)
        if len(ts) != len(res.P_in):
            raise ValueError("결과와 timestamps 길이가 다릅니다")
        res.timestamps = ts
        res.meta["n_years"] = len(ts.year.unique())
        res.meta.pop("degradation_fp", None)
        res.meta.pop("annual_table", None)
    return res


def _annual_n(res) -> float:
    return float(res.meta.get("n_years", len(res.P_in) * res.dt / HOURS_PER_YEAR))


def _annual_mean(res, value) -> float:
    return float(value) / max(_annual_n(res), 1e-12)


def annual_result_table(res: DispatchResult) -> pd.DataFrame:
    deg = ensure_degradation(res)
    cached = res.meta.get("annual_table")
    if cached is not None and res.meta.get("annual_table_fp") == res.meta["degradation_fp"]:
        return cached
    if res.timestamps is None:
        raise ValueError("연간/lifecycle 집계에는 dispatch 호출의 timestamps가 필요합니다")
    ts = pd.DatetimeIndex(res.timestamps)
    years, group = np.unique(ts.year, return_inverse=True)
    for y in years:
        expected = pd.date_range(f"{y}-01-01", f"{y+1}-01-01", freq=pd.Timedelta(hours=res.dt), inclusive="left")
        if not ts[ts.year == y].equals(expected):
            raise ValueError(f"{y}: 완전한 달력연도만 연평균/lifecycle에 사용할 수 있습니다")
    n = len(ts)
    if any(len(x) != n for x in (res.H2_series, res.P_used, res.j)):
        raise ValueError("연간 집계에는 keep_series=True 결과가 필요합니다")
    on_h = np.asarray(res.on_hours if res.on_hours is not None else (res.j > 0) * res.dt)
    paid = res.P_in * res.dt
    used = res.P_used * res.dt
    bop = res.H2_series * BOP_EFFICIENCY
    stack = used - bop
    cur = np.asarray(res.curtail_series if res.curtail_series is not None else np.zeros(n)) * res.dt
    idle = np.where((res.P_in > 0) & (on_h <= 0), paid, 0.0) if res.case == C1_NAME else np.zeros(n)
    dis = np.asarray(res.discharge if res.discharge is not None else np.zeros(n)) * res.dt
    eta = float(res.meta.get("eta_RTE", ETA_RTE))
    via = dis * eta
    direct = np.asarray(res.meta.get("direct_energy_series", used - via))
    stopped = res.dt - on_h
    values = dict(calendar_hours=np.full(n, res.dt), H2_kg=res.H2_series, op_hours=on_h,
                  idle_with_gen_h=np.where(res.P_in > 1e-9, stopped, 0),
                  no_gen_stopped_h=np.where(res.P_in <= 1e-9, stopped, 0),
                  E_paid_kWh=paid, E_used_kWh=used, E_stack_kWh=stack, E_bop_kWh=bop,
                  E_curtail_kWh=cur, E_idle_kWh=idle, E_rte_kWh=dis * (1 - eta),
                  E_direct_kWh=direct, E_via_ess_kWh=via,
                  eq_cycles=dis / res.E_rated if res.E_rated > 0 else np.zeros(n),
                  j_hours=np.asarray(res.j) * on_h)
    table = pd.DataFrame({k: np.bincount(group, weights=np.asarray(v, dtype=float), minlength=len(years))
                          for k, v in values.items()}, index=pd.Index(years, name="source_year"))
    si = np.asarray(res.meta.get("segment_input_index", np.arange(n)), dtype=int)
    event_group = group[si]
    for key, arr in [("degradation_V", deg["delta_uV"] * 1e-6),
                     ("starts", deg["start_flags"]), ("stops", deg["stop_flags"])]:
        table[key] = np.bincount(event_group, weights=np.asarray(arr, dtype=float), minlength=len(years))
    for key, arr in deg["contrib_uV_series"].items():
        table[f"degradation_{key}_V"] = np.bincount(event_group, weights=arr * 1e-6, minlength=len(years))
    table["E_unabsorbed_kWh"] = table.E_curtail_kWh + table.E_idle_kWh
    table["SEC_stack_kWh_kg"] = np.divide(table.E_stack_kWh, table.H2_kg,
                                         out=np.zeros(len(table)), where=table.H2_kg.to_numpy() > 0)
    table["rate_eff_uV_h"] = np.divide(table.degradation_V * 1e6, table.op_hours,
                                      out=np.zeros(len(table)), where=table.op_hours.to_numpy() > 0)
    table["j_mean_operating"] = np.divide(table.j_hours, table.op_hours,
                                         out=np.zeros(len(table)), where=table.op_hours.to_numpy() > 0)
    table = table.drop(columns="j_hours")
    end = np.r_[np.flatnonzero(np.diff(group)) + 1, n] - 1
    E_end = np.asarray(res.soc)[end] * res.E_rated if res.soc is not None else np.zeros(len(years))
    E_start = np.r_[res.E_start, E_end[:-1]]
    table["E_start_kWh"], table["E_end_kWh"] = E_start, E_end
    table["storage_delta_kWh"] = E_end - E_start
    table["ledger_residual_kWh"] = (table.E_paid_kWh - table.E_used_kWh - table.E_unabsorbed_kWh
                                     - table.E_rte_kWh - table.storage_delta_kWh)
    scale = np.maximum(table.E_paid_kWh.to_numpy(), 1)
    if np.any(abs(table.ledger_residual_kWh.to_numpy()) > 1e-7 * scale):
        raise RuntimeError(f"{res.case}: 연간 에너지 수지가 맞지 않습니다")
    res.meta["annual_table"], res.meta["annual_table_fp"] = table, res.meta["degradation_fp"]
    return table


def monthly_year_table(series, timestamps) -> pd.DataFrame:
    ts = pd.DatetimeIndex(timestamps)
    a = np.asarray(series, dtype=float)
    if len(a) != len(ts) or np.any(~np.isfinite(a)):
        raise ValueError("월별 집계의 시계열/시간축 길이 또는 유한성을 확인하세요")
    counts = pd.Series(1, index=ts).groupby([ts.year, ts.month]).sum().unstack()
    sums = pd.Series(a, index=ts).groupby([ts.year, ts.month]).sum().unstack()
    if counts.reindex(columns=range(1, 13)).isna().any().any():
        raise ValueError("월별 평균에는 모든 연도의 12개월이 필요합니다")
    return sums.reindex(columns=range(1, 13)).rename_axis(index="year", columns="month")


def monthly_statistics(series, timestamps) -> pd.DataFrame:
    table = monthly_year_table(series, timestamps)
    std = table.std(axis=0, ddof=1) if len(table) > 1 else table.iloc[0] * 0.0
    return pd.DataFrame({"mean": table.mean(axis=0), "std": std,
                         "min": table.min(axis=0), "max": table.max(axis=0), "n_years": len(table)})


# %%
# PART 9 — LCOH 경제성  +  스택 열화 · 교체 lifecycle 엔진
########################################################################################
########################################################################################

def ess_capex(E_rated_kWh: float, P_rated_kW: float,
              c_E: float | None = None, c_P: float | None = None) -> float:
    c_E = C_E_USD_KWH if c_E is None else c_E
    c_P = C_P_USD_KW if c_P is None else c_P
    return c_E * E_rated_kWh + c_P * P_rated_kW

def annual_fixed_om_components(capex_stack: float, capex_bop: float,
                               capex_ess: float = 0.0) -> tuple[float, float]:
    """PEM O&M = 5%, ESS O&M = 1.5% [USD/yr]."""
    cs, cb, ce = float(capex_stack), float(capex_bop), float(capex_ess)
    fp, fe = float(PEM_FIXED_OM_FRAC), float(ESS_FIXED_OM_FRAC)
    if not all(math.isfinite(v) and v >= 0.0 for v in (cs, cb, ce, fp, fe)):
        raise ValueError("CAPEX와 연간 고정 O&M 비율은 유한한 비음수여야 합니다")
    fixed_om_pem = (cs + cb) * fp
    fixed_om_ess = ce * fe
    return fixed_om_pem, fixed_om_ess

def stack_retention_at(age_h: float, rate_uV_h: float,
                       V_ref: float | None = None) -> float:
    """운전시간 age_h에서의 성능 유지율 (BOL = 1)."""
    age = max(float(age_h), 0.0)
    rate = max(float(rate_uV_h), 0.0)
    V0 = (float(np.atleast_1d(V(float(np.clip(DEG_J_REF, J_MIN, J_MAX))))[0])
          if V_ref is None else float(V_ref))
    dV = rate * 1.0e-6 * age
    return float(np.clip(V0 / max(V0 + dV, 1e-12), 1.0e-6, 1.0))

def h2_derate_factor(retention: float, sec_stack: float,
                     bop: float | None = None) -> float:
    """스택만 열화: H2 비율 = (SEC_stack + BOP) / (SEC_stack / retention + BOP)."""
    r = float(np.clip(retention, 1e-6, 1.0))
    b = BOP_EFFICIENCY if bop is None else float(bop)
    s = float(sec_stack)
    if not np.isfinite(s) or s <= 0:
        return r
    return (s + b) / (s / r + b)


# ── 9.2 lifecycle 스케줄 ──────────────────────────────────────────────────────

def lifecycle_schedule(annual_op_hours=None, annual_H2=None,
                       sec_stack_avg=float("nan"), rate_uV_h=0.0,
                       life_h=None, V_ref=None, years=None, *,
                       annual_inputs=None, eol_delta_V=None) -> pd.DataFrame:
    years = int(LIFETIME_YR if years is None else years)
    life_h = float(STACK_LIFETIME_H if life_h is None else life_h)
    V0 = float(V_ref if V_ref is not None else V(DEG_J_REF)[0])
    eol = float(eol_delta_V if eol_delta_V is not None else V0 * (1 / STACK_EOL_RETENTION - 1))
    if years < 1 or life_h <= 0 or not np.isfinite(life_h) or V0 <= 0 or eol <= 0:
        raise ValueError("프로젝트 연수, 스택 수명, 기준전압, EOL 손상 한계는 양수여야 합니다")
    if annual_inputs is None:
        annual_inputs = pd.DataFrame([dict(source_year=0, calendar_hours=HOURS_PER_YEAR,
            op_hours=float(annual_op_hours), H2_kg=float(annual_H2),
            SEC_stack_kWh_kg=sec_stack_avg, rate_eff_uV_h=rate_uV_h)]).set_index("source_year")
    if not len(annual_inputs):
        raise ValueError("연도별 입력 결과가 없습니다")
    age = damage = pending = 0.0
    rows, tol = [], 1e-9
    for year in range(1, years + 1):
        a = annual_inputs.iloc[(year - 1) % len(annual_inputs)]
        source_year = int(annual_inputs.index[(year - 1) % len(annual_inputs)])
        hours, potential, bol_h2 = float(a.calendar_hours), float(a.op_hours), float(a.H2_kg)
        rate = max(float(a.rate_eff_uV_h), 0.0) * 1e-6
        sec = float(a.SEC_stack_kWh_kg)
        if not np.isfinite(hours + potential + bol_h2 + rate) or hours <= 0 or not 0 <= potential <= hours + tol or bol_h2 < 0:
            raise ValueError(f"{source_year}: 연간 운전/생산 입력이 유효하지 않습니다")
        duty = potential / hours
        remaining = hours
        h2 = run_total = ret_weighted = lost = lost_calendar = 0.0
        age_start, damage_start, pending_start = age, damage, pending
        n_repl = 0
        for _ in range(100_000):
            if remaining <= tol:
                break
            if pending > tol:
                if DOWNTIME_DUTY_WEIGHTED:
                    take = min(remaining, pending)
                    pending -= take
                elif duty > 0:
                    take = min(remaining, pending / duty)
                    pending -= take * duty
                else:
                    take = remaining
                remaining -= take
                lost += take * duty
                lost_calendar += take
                if remaining <= tol:
                    break
            if duty <= 0:
                remaining = 0.0
                break
            to_voltage = max(eol - damage, 0) / rate if rate > 0 else float("inf")
            run = min(remaining * duty, max(life_h - age, 0), to_voltage)
            if run > tol:
                increment = rate * run
                # 구간 평균 유지율 및 H2 derating의 해석적 적분.
                ret = V0 / (V0 + damage)
                alpha = (sec / ((sec + BOP_EFFICIENCY) * V0)
                         if np.isfinite(sec) and sec > 0 else 1.0 / V0)
                factor = 1.0 / (1.0 + alpha * damage)
                if increment > 1e-14:
                    ret = V0 * math.log1p(increment / (V0 + damage)) / increment
                    factor = math.log1p(alpha * increment / (1 + alpha * damage)) / (alpha * increment)
                h2 += bol_h2 * run / potential * factor
                ret_weighted += ret * run
                age += run
                damage += increment
                run_total += run
                remaining = max(remaining - run / duty, 0.0)
            at_eol = age >= life_h - tol or damage >= eol - 1e-12
            if at_eol:
                if remaining <= tol and year == years:
                    break
                n_repl += 1
                age = damage = 0.0
                pending += max(float(REPLACEMENT_DOWNTIME_H), 0)
            elif run <= tol:
                raise RuntimeError("lifecycle의 시간 적분이 진행하지 않습니다")
        else:
            raise RuntimeError("연간 교체 이벤트가 과도합니다")
        uptime = run_total / potential if potential > 0 else 0.0
        rows.append(dict(year=year, source_year=source_year, source_cycle=(year - 1) // len(annual_inputs) + 1,
            calendar_hours=hours, BOL_H2_kg=bol_h2, BOL_op_hours=potential,
            op_hours=run_total, uptime=uptime, H2_kg=h2,
            retention_avg=ret_weighted / run_total if run_total > 0 else V0 / (V0 + damage),
            retention_end=V0 / (V0 + damage),
            derate=h2 / (bol_h2 * uptime) if bol_h2 * uptime > 0 else 1.0,
            n_replacement=n_repl, age_start_h=age_start, age_end_h=age,
            damage_start_V=damage_start, damage_end_V=damage,
            downtime_lost_op_h=lost, downtime_lost_calendar_h=lost_calendar,
            downtime_pending_op_h=max(pending, 0) * (duty if DOWNTIME_DUTY_WEIGHTED else 1),
            downtime_pending_h=max(pending, 0), downtime_pending_start_h=pending_start))
    return pd.DataFrame(rows)



# ── 9.3 lifecycle LCOH (열화 · 교체 반영 DCF) ─────────────────────────────────

_LIFECYCLE_COST_KEYS = ("electricity", "curtailment_penalty", "fixed_OM",
                        "variable_OM", "water", "stack_replacement",
                        "battery_replacement")


def lifecycle_lcoh(res: DispatchResult, stack: FixedStack, capex_ess=0.0,
                   elec_price=None, penalty_mult=None, years=None) -> dict:
    """관측연도별 BOL 결과를 사업기간에 순환 적용하는 연간 집계 DCF."""
    price = float(ELEC_PRICE if elec_price is None else elec_price)
    mult = float(ELEC_PENALTY_MULT if penalty_mult is None else penalty_mult)
    years = int(LIFETIME_YR if years is None else years)
    annual = annual_result_table(res)
    annual_H2, op_hours = float(annual.H2_kg.mean()), float(annual.op_hours.mean())
    capex_total = stack.capex_stack + stack.capex_bop + capex_ess
    fixed_om_pem, fixed_om_ess = annual_fixed_om_components(stack.capex_stack, stack.capex_bop, capex_ess)
    fixed_om = fixed_om_pem + fixed_om_ess
    deg = ensure_degradation(res)
    V_ref = float(V(deg["j_mean_operating"])[0])
    sched = lifecycle_schedule(annual_inputs=annual, years=years, life_h=STACK_LIFETIME_H,
                               V_ref=V_ref, eol_delta_V=deg["eol_delta_V"])
    rows = []
    for s in sched.to_dict("records"):
        a = annual.loc[int(s["source_year"])]
        up, h2 = float(s["uptime"]), float(s["H2_kg"])
        paid = float(a.E_paid_kWh) * (1.0 if ELEC_PAID_DURING_DOWNTIME else up)
        unabsorbed = (float(a.E_unabsorbed_kWh) + float(a.E_used_kWh) * (1 - up)
                      if ELEC_PAID_DURING_DOWNTIME else float(a.E_unabsorbed_kWh) * up)
        rows.append(dict(s, E_used_kWh=float(a.E_used_kWh) * up, E_paid_kWh=paid,
            electricity=paid * price, curtailment_penalty=price * (mult - 1) * unabsorbed,
            fixed_OM_pem=fixed_om_pem, fixed_OM_ess=fixed_om_ess, fixed_OM=fixed_om,
            variable_OM=h2 * VARIABLE_OM, water=h2 * WATER_STOICH * WATER_PRICE,
            stack_replacement=int(s["n_replacement"]) * stack.capex_stack * REPL_STACK_FRAC,
            battery_replacement=ANNUAL_BATT_REPL_USD))
    life = pd.DataFrame(rows)
    life["opex"] = life[list(_LIFECYCLE_COST_KEYS)].sum(axis=1)
    den = life.H2_kg.replace(0, np.nan)
    life["SEC_used_kWh_kg"] = life.E_used_kWh / den
    life["SEC_eff_kWh_kg"] = life.E_paid_kWh / den
    life["eff_HHV_pct"] = HHV_H2_KWH_KG / life.SEC_used_kWh_kg * 100
    life["eff_HHV_paid_pct"] = HHV_H2_KWH_KG / life.SEC_eff_kWh_kg * 100
    disc = (1 + INTEREST_RATE) ** life.year.to_numpy(dtype=float)
    life["discount_factor"] = 1 / disc
    life["discounted_opex"] = life.opex / disc
    life["discounted_H2_kg"] = life.H2_kg / disc
    pv_factor = float(np.sum(1 / disc))
    pv_h2 = float(life.discounted_H2_kg.sum())
    pv_opex = {k: float((life[k] / disc).sum()) for k in _LIFECYCLE_COST_KEYS}
    pv_cost = capex_total + sum(pv_opex.values())
    lcoh = pv_cost / pv_h2 if pv_h2 > 0 else float("inf")
    items_usd = dict(capex=capex_total, **pv_opex)
    items_per_kg = {k: v / pv_h2 for k, v in items_usd.items()} if pv_h2 > 0 else {}
    ess_pv = capex_ess + (fixed_om_ess + ANNUAL_BATT_REPL_USD) * pv_factor
    return dict(LCOH=lcoh, annual_cost=pv_cost / pv_factor, annual_H2=annual_H2,
        annual_H2_std=float(annual.H2_kg.std(ddof=1)) if len(annual) > 1 else 0.0,
        annual_H2_mean=float(life.H2_kg.mean()), H2_total_life_kg=float(life.H2_kg.sum()),
        items_usd=items_usd, items_usd_per_kg=items_per_kg,
        fixed_OM_pem_annual_usd=fixed_om_pem, fixed_OM_ess_annual_usd=fixed_om_ess,
        capex_stack=stack.capex_stack, capex_bop=stack.capex_bop, capex_ess=capex_ess,
        ess_annual_usd=ess_pv / pv_factor, ess_share_pct=ess_pv / pv_cost * 100 if pv_cost else 0,
        op_hours=op_hours, lifecycle=life, annual_inputs=annual, project_years=years,
        replacement_count=int(life.n_replacement.sum()),
        retention_mean_pct=float(np.average(life.retention_avg, weights=life.op_hours)) * 100 if life.op_hours.sum() > 0 else 100,
        retention_end_pct=float(life.retention_end.iloc[-1] * 100),
        H2_fade_pct=float((1 - life.derate.iloc[-1]) * 100),
        PV_H2_kg=pv_h2, PV_cost_USD=pv_cost, pv_factor=pv_factor,
        H2_mean_vs_BOL_pct=float(life.H2_kg.sum() / max(life.BOL_H2_kg.sum(), 1e-12) * 100),
        stack_life_hours=deg["stack_life_hours"], stack_life_years=deg["stack_life_years"],
        stack_replacements_project=int(life.n_replacement.sum()), degradation=deg,
        source_years=tuple(int(y) for y in annual.index),
        basis="multiyear lifecycle (연도별 BOL 순환, 열화·교체 반영 DCF)")


def lcoh_compare(res: DispatchResult, stack: FixedStack, capex_ess: float = 0.0,
                 elec_price: float | None = None,
                 penalty_mult: float | None = None) -> dict:
    mult = ELEC_PENALTY_MULT if penalty_mult is None else penalty_mult
    out = lifecycle_lcoh(res, stack, capex_ess, elec_price, mult)
    out["accounting"] = "allocated"
    out["basis"] = "multiyear lifecycle (연도별 BOL 순환, 열화·교체 반영 DCF, 배정 kWh)"
    out["penalty_mult"] = mult
    out["E_unabsorbed"] = _annual_mean(res, res.E_unabsorbed)
    return out

def lcoh_internal(res: DispatchResult, stack: FixedStack,
                  capex_ess: float = 0.0, elec_price: float | None = None,
                  penalty_mult: float | None = None) -> float:
    return float(lcoh_compare(res, stack, capex_ess, elec_price, penalty_mult)["LCOH"])

def report_lcoh(name: str, out: dict) -> None:
    items = out["items_usd_per_kg"]
    print(f"\n[{name}] LCOH={out['LCOH']:.4f} USD/kg-H2 | "
          f"H2={out['annual_H2']:,.0f} kg/yr (BOL) | "
          f"accounting={out.get('accounting', 'allocated')}")
    print(f"  basis: {out.get('basis', '-')}")
    print(f"  CAPEX stack/BOP/ESS = {out['capex_stack']:,.0f} / "
          f"{out['capex_bop']:,.0f} / {out['capex_ess']:,.0f} USD, "
          f"ESS annual share={out['ess_share_pct']:.2f} %")
    if "fixed_OM_pem_annual_usd" in out and "fixed_OM_ess_annual_usd" in out:
        om_pem = out["fixed_OM_pem_annual_usd"]
        om_ess = out["fixed_OM_ess_annual_usd"]
        print(f"  Fixed O&M [USD/yr]: PEM(stack+BOP)={om_pem:,.2f}, "
              f"ESS={om_ess:,.2f}, total={om_pem + om_ess:,.2f}")
    print("  USD/kg: " + ", ".join(f"{k}={v:.3f}" for k, v in items.items()))
    if "lifecycle" in out and len(out.get("lifecycle", [])):
        life = out["lifecycle"]
        print(f"  lifecycle: 교체 {out['replacement_count']}회, "
              f"생애 H2={out['H2_total_life_kg'] / 1e3:,.1f} t "
              f"({LIFETIME_YR} yr), 최종연도 열화손실={out['H2_fade_pct']:.2f} %, "
              f"평균 유지율={out['retention_mean_pct']:.2f} %")
        print(f"  system efficiency(HHV): "
              f"{life['eff_HHV_pct'].iloc[0]:.2f} % (yr1) -> "
              f"{life['eff_HHV_pct'].iloc[-1]:.2f} % (yr{LIFETIME_YR})")


# %%
# PART 10 — 그래프 공용 유틸
########################################################################################
########################################################################################

CASE_COLORS = {"Case 1": "#1f6f8b", "Case 2": "#2a9d8f", "Case 3": "#d1495b"}
LEDGER_COLORS = {"curtail": "#e76f51", "idle": "#9c6644", "rte": "#e9c46a",
                 "bop": "#a8b8c8", "stack": "#2a9d8f", "storage_delta": "#808080"}


def setup_plots() -> None:
    if matplotlib is None:
        return
    matplotlib.rcParams.update({
        "figure.dpi": FIG_DPI, "axes.grid": True, "grid.alpha": 0.25,
        "axes.spines.top": False, "axes.spines.right": False, "font.size": 10,
        "axes.titlesize": 12, "axes.labelsize": 10, "legend.frameon": False})
    for cand in ("NanumGothic", "Malgun Gothic", "AppleGothic", "Noto Sans CJK KR"):
        try:
            matplotlib.font_manager.findfont(cand, fallback_to_default=False)
            matplotlib.rcParams["font.family"] = cand
            matplotlib.rcParams["axes.unicode_minus"] = False
            break
        except Exception:
            continue

def _finish(fig, name: str) -> None:
    import os
    _FIGURE_MANIFEST.append({"filename": f"{name}.png", "titles": " | ".join(ax.get_title() for ax in fig.axes)})
    d = _ensure_fig_dir()
    if d:
        fig.savefig(os.path.join(d, f"{name}.png"), bbox_inches="tight", dpi=FIG_DPI)
    if SHOW_FIGURES:
        plt.show()
    else:
        plt.close(fig)



# %%
# PART 11 — CASE 1 : 직결 (Direct coupling)
########################################################################################
########################################################################################

C1_NAME = "Case 1"


def dispatch_case1(P_in: np.ndarray, A_tot: float, j_min: float, j_max: float,
                   dt: float = 1.0, *, timestamps=None, initial_j=None,
                   warm_start=None) -> DispatchResult:
    P = np.asarray(P_in, dtype=float)
    P_min = float(P_of_j(j_min, A_tot)[0])
    P_max = float(P_of_j(j_max, A_tot)[0])
    inv = get_inverter(A_tot, j_min, j_max)     # inv = P -> j 계산한 거 찾기

    m_above = P > P_max
    m_idle = (P > 0.0) & (P < P_min)
    m_night = P <= 0.0
    m_in = ~(m_above | m_idle | m_night)

    j = np.zeros_like(P)
    j[m_above] = j_max
    if np.any(m_in):
        j[m_in] = inv(P[m_in])
    on = j > 0.0
    enabled = bool(WARM_START if warm_start is None else warm_start)
    previous_j = (float(j[-1]) if enabled and len(j) and WARM_START_DAYS > 0 else 0.0) if initial_j is None else float(initial_j)

    h2_rate = np.zeros_like(P)                      # [kg/h]
    h2_rate[on] = h2_area(j[on]) * A_tot
    H2_series = h2_rate * dt                        # [kg]  rate/amt 분리 (Δt 대비)

    P_stack = np.zeros_like(P)
    P_stack[on] = j[on] * V(j[on]) * A_tot / 1000.0
    P_bop = BOP_EFFICIENCY * h2_rate
    P_used = P_stack + P_bop

    E_paid = float(P.sum() * dt)                            # 배정 전량 지불 (R7)
    E_curtail = float((P[m_above] - P_max).sum() * dt)      # 초과분만 버림
    E_idle = float(P[m_idle].sum() * dt)                    # 하한 미만 -> 전량 손실
    E_stack = float(P_stack.sum() * dt)
    E_bop = float(P_bop.sum() * dt)

    cur = np.zeros_like(P)
    cur[m_above] = P[m_above] - P_max

    result = DispatchResult(
        case=C1_NAME, dt=dt, P_in=P, j=j, P_used=P_used, H2_series=H2_series,
        curtail_series=cur,
        E_paid=E_paid, E_curtail=E_curtail, E_idle=E_idle, E_rte=0.0,
        E_bop=E_bop, E_stack=E_stack, E_direct=E_stack + E_bop, E_via_ess=0.0,
        H2_total=float(H2_series.sum()), op_hours=float(on.sum() * dt),
        on_hours=on.astype(float) * dt,
        restarts=int(np.sum(on & ~np.r_[previous_j > 0,
                                       on[:-1]])),
        hours_curtail=float(m_above.sum() * dt), hours_idle=float(m_idle.sum() * dt),
        hours_night=float(m_night.sum() * dt),
        meta={"P_min": P_min, "P_max": P_max, "j_min": j_min, "j_max": j_max,
              "initial_j": previous_j, "warm_start_excluded": True,
              "warm_start_scope": "once_before_full_analysis",
              "warm_start_passes": int(enabled and WARM_START_DAYS > 0),
              "warm_start_hours": min(len(P) * dt, WARM_START_DAYS * 24) if enabled else 0.0})

    return _bind_result(result, timestamps)



def run_case1(profile=None, fixed_stack=None, j_star=None, verbose=True) -> dict:
    profile = profile or get_profile()
    fixed_stack = fixed_stack or build_fixed_stack(profile, verbose=verbose)
    win = operating_window(fixed_stack.A_tot)

    if verbose:
        print(f"\n[PART 11] {C1_NAME} — direct coupling")
        print(win.describe())

    res = dispatch_case1(profile.P_in, fixed_stack.A_tot, win.j_min, win.j_max,
                         dt=profile.dt, timestamps=profile.timestamps)
    lc = lcoh_compare(res, fixed_stack, 0.0)          # ESS 없음

    if verbose:
        res.report(h2_area_at_rated=float(h2_area(win.j_max)[0]),
                   h2_area_at_jstar=float(h2_area(j_star)[0]) if j_star else None,
                   A_tot=fixed_stack.A_tot)
        _report_case1_extra(res, win)
        report_lcoh(C1_NAME, lc)

    return {"result": res, "lcoh": lc, "window": win,
            "stack": fixed_stack, "profile": profile, "capex_ess": 0.0}


def _report_case1_extra(res: DispatchResult, win: OperatingWindow) -> None:
    on = res.j > 0
    print("\n[Case 1 운전 요약]")
    print(f"  curtail={res.hours_curtail:,.0f} h ({res.E_curtail / res.E_paid * 100:.2f} %), "
          f"idle={res.hours_idle:,.0f} h ({res.E_idle / res.E_paid * 100:.2f} %), "
          f"night={res.hours_night:,.0f} h")
    if on.any():
        w = res.H2_series[on]
        jw = float((res.j[on] * w).sum() / w.sum()) if w.sum() > 0 else np.nan
        print(f"  j[min/med/mean/max]={res.j[on].min():.4f}/"
              f"{np.median(res.j[on]):.4f}/{res.j[on].mean():.4f}/"
              f"{res.j[on].max():.4f} A/cm2 | H2-weighted j={jw:.4f}")
    print(f"  SEC_stack_avg={res.SEC_stack_avg():.2f}, "
          f"rated SEC_stack={float(SEC_stack(win.j_max)[0]):.2f} kWh/kg")


def stack_rating_sweep_case1(profile=None, factors=None, base_nameplate=None,
                             verbose=True) -> pd.DataFrame:
    profile = profile or get_profile()
    factors = factors or CASE1["rating_sweep_factors"]
    if base_nameplate is None:
        base_nameplate = build_fixed_stack(profile, verbose=False).P_nameplate

    rows = []
    for f in factors:
        st = size_fixed_stack(base_nameplate * f)
        w = operating_window(st.A_tot)
        r = dispatch_case1(profile.P_in, st.A_tot, w.j_min, w.j_max, dt=profile.dt, timestamps=profile.timestamps)
        lc = lcoh_compare(r, st, 0.0)
        rows.append({"nameplate [kW]": st.P_nameplate, "N_cells": st.N_cells,
                     "curtail [%]": r.E_curtail / r.E_paid * 100,
                     "idle [%]": r.E_idle / r.E_paid * 100,
                     "op_hours [h/yr]": _annual_mean(r, r.op_hours), "annual_H2 [kg]": _annual_mean(r, r.H2_total),
                     "SEC_eff [kWh/kg]": r.SEC_eff(),
                     "CAPEX_fix [kUSD]": st.capex_fix / 1e3, "LCOH [$/kg]": lc["LCOH"]})
    df = pd.DataFrame(rows)
    if verbose:
        k = int(df["LCOH [$/kg]"].idxmin())
        print(f"\n[Case 1 rating sweep] best={df.loc[k, 'nameplate [kW]']:,.0f} kW, "
              f"LCOH={df.loc[k, 'LCOH [$/kg]']:.3f}, "
              f"curtail={df.loc[k, 'curtail [%]']:.2f} %")
    return df


# %%
# PART 12 — CASE 2 : ESS 완충 (ESS-buffered, 고정 j*)
########################################################################################
########################################################################################

C2_NAME = "Case 2"


def c2_operating_point(j: float, A_tot: float) -> dict:
    P_stack, P_bop = split_power(j, A_tot)
    return {"j": float(j), "P_star": P_stack + P_bop, "P_stack": P_stack,
            "P_bop": P_bop, "h2_rate": float(h2_area(j)[0]) * A_tot}


def c2_charge_rating(profile, P_star: float) -> float:
    if _use_charge_pctl():
        return max(profile.pctl(CHARGE_RATING_PCTL), P_star)
    return profile.peak


class _StorageTrace:
    def __init__(self, P, op, E_rated, E_start, eta, dt, keep_series, initial_running):
        self.P, self.op = P, op
        self.E_rated, self.E_start, self.eta, self.dt = E_rated, E_start, eta, dt
        self.keep = bool(keep_series)
        n = len(P)
        self.on = np.zeros(n)
        self.chg = np.zeros(n)
        self.dis = np.zeros(n)
        self.direct = np.zeros(n)
        self.curtail = np.zeros(n)
        self.curtail_h = np.zeros(n)
        self.soc = np.zeros(n)
        self.segment_j, self.segment_h, self.segment_i = [], [], []
        self.segment_chg, self.segment_dis, self.segment_E = [], [], []
        self.initial_j = float(op["j"]) if initial_running else 0.0
        self.last_j = self.initial_j
        self.starts = self.stops = 0

    def record(self, i, hours, running, charge_kw, draw_kw, direct_kw, energy):
        j = float(self.op["j"]) if running else 0.0
        if j > 0 and self.last_j <= 0:
            self.starts += 1
        if j <= 0 and self.last_j > 0:
            self.stops += 1
        changed = j != self.last_j
        self.last_j = j
        if hours > 0:
            cur = max(float(self.P[i]) - charge_kw - direct_kw, 0.0)
            self.on[i] += hours if running else 0.0
            self.chg[i] += charge_kw * hours
            self.dis[i] += draw_kw * hours
            self.direct[i] += direct_kw * hours
            self.curtail[i] += cur * hours
            self.curtail_h[i] += hours if cur > 1e-9 else 0.0
        if self.keep and (hours > 0 or changed):
            self.segment_j.append(j)
            self.segment_h.append(float(hours))
            self.segment_i.append(i)
            self.segment_chg.append(charge_kw)
            self.segment_dis.append(draw_kw)
            self.segment_E.append(energy)

    def stop(self, i, energy):
        self.record(i, 0.0, False, 0.0, 0.0, 0.0, energy)

    def finish(self, case, E_end, P_rc, P_rd, running_end, meta):
        on_h = float(self.on.sum())
        h2 = self.on * self.op["h2_rate"]
        used = self.on * self.op["P_star"] / self.dt
        j = np.where(self.on > 1e-12, self.op["j"], 0.0)
        info = dict(meta, initial_j=self.initial_j, running_end=bool(running_end),
                    initial_running=bool(self.initial_j > 0), stops=self.stops,
                    generation_absent_hours=float((self.P <= 1e-9).sum() * self.dt),
                    j_series_semantics="actual ON current; use on_hours for duration")
        if self.keep:
            info.update(direct_energy_series=self.direct.copy(),
                        curtail_hours_series=self.curtail_h.copy(),
                        segment_j=np.asarray(self.segment_j),
                        segment_duration_h=np.asarray(self.segment_h),
                        segment_input_index=np.asarray(self.segment_i, dtype=int),
                        segment_charge_kW=np.asarray(self.segment_chg),
                        segment_discharge_kW=np.asarray(self.segment_dis),
                        segment_E_end_kWh=np.asarray(self.segment_E))
        return DispatchResult(
            case=case, dt=self.dt, P_in=self.P,
            j=j if self.keep else np.zeros(0),
            P_used=used if self.keep else np.zeros(0),
            H2_series=h2 if self.keep else np.zeros(0),
            on_hours=self.on, soc=self.soc if self.keep else None,
            running=(self.on > 1e-12) if self.keep else None,
            charge=self.chg / self.dt if self.keep else None,
            discharge=self.dis / self.dt if self.keep else None,
            curtail_series=self.curtail / self.dt if self.keep else None,
            E_paid=float(self.P.sum() * self.dt), E_curtail=float(self.curtail.sum()),
            E_idle=0.0, E_rte=float(self.dis.sum() * (1.0 - self.eta)),
            E_bop=on_h * self.op["P_bop"], E_stack=on_h * self.op["P_stack"],
            E_direct=float(self.direct.sum()), E_via_ess=float(self.dis.sum() * self.eta),
            H2_total=float(h2.sum()), op_hours=on_h,
            hours_curtail=float(self.curtail_h.sum()),
            hours_idle=float((self.dt - self.on)[self.P > 1e-9].sum()),
            hours_night=float((self.dt - self.on)[self.P <= 1e-9].sum()),
            restarts=self.starts, eq_cycles=float(self.dis.sum() / self.E_rated) if self.E_rated > 0 else 0.0,
            E_start=self.E_start, E_end=E_end, E_rated=self.E_rated,
            P_rated_chg=P_rc, P_rated_dis=P_rd,
            j_star=self.op["j"], P_star=self.op["P_star"], meta=info)


def _storage_inputs(P_in, op, E_rated, P_rc, P_rd, eta, dt, soc_max, soc_floor, soc0):
    P = np.asarray(P_in, dtype=float)
    if P.ndim != 1 or not len(P) or np.any(~np.isfinite(P)) or np.any(P < 0):
        raise ValueError("P_in은 유한한 비음수 1차원 배열이어야 합니다")
    if not 0 < eta <= 1 or not np.isfinite(dt) or dt <= 0:
        raise ValueError("0 < eta_RTE <= 1, dt > 0이 필요합니다")
    if not 0 <= soc_floor < soc_max <= 1:
        raise ValueError("0 <= soc_floor < soc_max <= 1이 필요합니다")
    if any(not np.isfinite(v) or v < 0 for v in (E_rated, P_rc, P_rd)):
        raise ValueError("ESS 에너지/출력 정격은 유한한 비음수여야 합니다")
    if any(not np.isfinite(op[k]) or op[k] <= 0 for k in ("P_star", "j", "h2_rate")):
        raise ValueError("운전점 전력/전류/생산율이 양수여야 합니다")
    if abs(op["P_stack"] + op["P_bop"] - op["P_star"]) > max(op["P_star"], 1) * 1e-9:
        raise ValueError("운전점에서 stack + BOP != P_star")
    s0 = soc_floor if soc0 is None else float(soc0)
    if not np.isfinite(s0) or s0 < soc_floor - 1e-10 or s0 > soc_max + 1e-10:
        raise ValueError("초기 SOC가 허용 범위 밖입니다")
    E = float(np.clip(s0, soc_floor, soc_max)) * E_rated
    return P, E


def _run_storage_warmstart(dispatch, args, E_rated, dt, kwargs):
    kw = dict(kwargs)
    timestamps = kw.pop("timestamps", None)
    requested = kw.pop("warm_start", None)
    enabled = bool(WARM_START if requested is None else requested)
    keep = bool(kw.pop("keep_series", True))
    n_warm = min(len(args[0]), int(round(float(WARM_START_DAYS) * 24 / dt))) if enabled else 0
    if n_warm < 0:
        raise ValueError("WARM_START_DAYS는 비음수여야 합니다")
    if n_warm and E_rated > 0:
        spin_args = (np.asarray(args[0])[-n_warm:], *args[1:])
        first = dispatch(*spin_args, dt=dt, keep_series=False, **kw)
        kw.update(soc0=first.E_end / E_rated, running0=bool(first.meta["running_end"]))
        if "controller_end" in first.meta:
            kw["controller_state"] = dict(first.meta["controller_end"])
    else:
        n_warm = 0
    r = dispatch(*args, dt=dt, keep_series=keep, **kw)
    r.meta.update(warm_start_passes=int(n_warm > 0), warm_start_hours=n_warm * dt,
                  warm_start_excluded=True, warm_start_scope="once_before_full_analysis",
                  warm_start_delta_kWh=r.E_end - r.E_start)
    if n_warm and timestamps is not None:
        r.meta.update(warm_start_source_start=str(timestamps[-n_warm]),
                      warm_start_source_end=str(timestamps[-1]))
    return _bind_result(r, timestamps)



def dispatch_case2(P_in, op, A_tot, E_rated, P_rated_chg,
                   soc_max=None, soc_floor=None, eta_RTE=None, soc0=None, dt=1.0,
                   keep_series=True, running0=False, control_interval_h=None,
                   charge_while_running=None, controller_state=None) -> DispatchResult:
    hi = float(CASE2["soc_max"] if soc_max is None else soc_max)
    lo = float(CASE2["soc_floor"] if soc_floor is None else soc_floor)
    eta = float(ETA_RTE if eta_RTE is None else eta_RTE)
    period = float(CASE2["control_interval_h"] if control_interval_h is None else control_interval_h)
    charge_on = bool(CASE2["charge_while_running"] if charge_while_running is None else charge_while_running)
    if not np.isfinite(period) or period <= 0:
        raise ValueError("Case 2 control_interval_h는 양수여야 합니다")
    P, E = _storage_inputs(P_in, op, E_rated, P_rated_chg, op["P_star"], eta,
                           dt, hi, lo, soc0)
    E_hi, E_lo = E_rated * hi, E_rated * lo
    Etol, Ttol = max(1e-10, E_rated * 1e-12), 1e-12
    draw = op["P_star"] / eta
    running = bool(running0)
    clock = dict(controller_state or {"remaining_h": 0.0})
    until_control = float(clock["remaining_h"])
    if not 0 <= until_control <= period + Ttol:
        raise ValueError("Case 2 controller_state가 유효하지 않습니다")
    tr = _StorageTrace(P, op, E_rated, E, eta, dt, keep_series, running)
    for i, p in enumerate(P):
        remaining = float(dt)
        for _ in range(100_000):
            if remaining <= Ttol:
                break
            if running and E <= E_lo + Etol:
                running = False
                tr.stop(i, E)
            if until_control <= Ttol:
                until_control = period
                if not running and E > E_lo + Etol:
                    running = True
            q = min(float(p), P_rated_chg) if (not running or charge_on) else 0.0
            d = draw if running else 0.0
            if E >= E_hi - Etol:
                q = min(q, d)
            net = q - d
            h = min(remaining, until_control)
            if net > 0:
                h = min(h, max(E_hi - E, 0.0) / net)
            elif net < 0:
                h = min(h, max(E - E_lo, 0.0) / (-net))
            if h <= Ttol:
                raise RuntimeError("Case 2 시간 적분이 진행하지 않습니다. SOC/제어주기를 확인하십시오")
            E += net * h
            remaining = max(remaining - h, 0.0)
            until_control = max(until_control - h, 0.0)
            tr.record(i, h, running, q, d, 0.0, E)
            if E < E_lo - 10 * Etol or E > E_hi + 10 * Etol:
                raise RuntimeError("Case 2 SOC 수지 위반")
            if running and E <= E_lo + Etol:
                running = False
                tr.stop(i, E)
        else:
            raise RuntimeError("Case 2 한 입력 구간의 이벤트 수가 과도합니다")
        tr.soc[i] = E / E_rated if E_rated > 0 else 0.0
    final_clock = {"remaining_h": 0.0 if until_control <= Ttol else until_control}
    return tr.finish(C2_NAME, E, P_rated_chg, op["P_star"], running,
                     {"E_hi": E_hi, "E_lo": E_lo, "eta_RTE": eta,
                      "control_interval_h": period, "charge_while_running": charge_on,
                      "storage_first": True, "controller_start": clock,
                      "controller_end": final_clock})


def c2_run_warmstart(P_in, op, A_tot, E_rated, P_rc, dt=1.0, **kw):
    return _run_storage_warmstart(dispatch_case2, (P_in, op, A_tot, E_rated, P_rc),
                                  E_rated, dt, kw)


def find_j_star_case2(profile, fixed_stack, cap_hours, n_j=None, verbose=True):
    n_j = n_j or _n_j_grid("case2")
    A_tot = fixed_stack.A_tot
    rows, best = [], None
    for j in np.linspace(J_MIN, J_MAX, n_j):
        op = c2_operating_point(j, A_tot)
        E_rated = cap_hours * op["P_star"]
        P_rc = c2_charge_rating(profile, op["P_star"])
        r = c2_run_warmstart(profile.P_in, op, A_tot, E_rated, P_rc,
                     dt=profile.dt, timestamps=profile.timestamps)
        if r.H2_total <= 0:
            continue
        cx = ess_capex(E_rated, max(P_rc, op["P_star"]))
        li = lcoh_internal(r, fixed_stack, cx)
        rows.append({"j": j, "P_star": op["P_star"], "E_rated": E_rated,
                     "op_hours": _annual_mean(r, r.op_hours),
                     "curtail_%": r.E_curtail / r.E_paid * 100,
                     "annual_H2_t": _annual_mean(r, r.H2_total) / 1000, "LCOH_internal": li,
                     "restarts": _annual_mean(r, r.restarts)})
        if best is None or li < best[2]:
            best = (j, op["P_star"], li, r)

    df = pd.DataFrame(rows)
    if verbose and len(df):
        k = int(df["LCOH_internal"].idxmin())
        print(f"\n[Case 2 j* search] cap={cap_hours:g} h, grid={n_j}, "
              f"j*={df.loc[k, 'j']:.4f} A/cm2, "
              f"P*={df.loc[k, 'P_star']:,.1f} kW, "
              f"LCOH_int={df.loc[k, 'LCOH_internal']:.3f}")
    return best, df


@dataclass(frozen=True)
class ESSDesign:
    case: str
    j_star: float
    E_rated: float
    P_rated_chg: float
    P_rated_dis: float
    sizing_q: float | None = None


class NoFeasibleDesignError(RuntimeError):
    pass


def evaluate_fixed_design(profile, fixed_stack, design: ESSDesign, *,
                          warm_start=None, initial_state=None) -> dict:
    """설비를 다시 사이징하지 않고 고정 설계를 평가."""
    if design.case not in (C2_NAME, C3_NAME):
        raise ValueError("ESSDesign.case는 'Case 2' 또는 'Case 3'여야 합니다")
    if not J_MIN <= design.j_star <= J_MAX:
        raise ValueError("고정 운전 전류밀도가 운전 창 밖입니다")
    op = c2_operating_point(design.j_star, fixed_stack.A_tot)
    kwargs = dict(initial_state or {})
    kwargs.update(timestamps=profile.timestamps, warm_start=warm_start)
    if design.case == C2_NAME:
        if design.P_rated_dis + 1e-8 < op["P_star"]:
            raise ValueError("Case 2의 P*가 고정 ESS 방전 정격보다 큽니다")
        r = c2_run_warmstart(profile.P_in, op, fixed_stack.A_tot, design.E_rated,
                            design.P_rated_chg, dt=profile.dt, **kwargs)
        r.P_rated_dis = design.P_rated_dis
    else:
        r = c3_run_warmstart(profile.P_in, op, fixed_stack.A_tot, design.E_rated,
                            design.P_rated_chg, design.P_rated_dis, dt=profile.dt,
                            restart_hours=CASE3["restart_hours_of_Pstar"], **kwargs)
    cx = ess_capex(design.E_rated, max(design.P_rated_chg, design.P_rated_dis))
    lc = lcoh_compare(r, fixed_stack, cx)
    return dict(result=r, lcoh=lc, op=op, design=design, j_star=design.j_star,
                cap_hours=design.E_rated / op["P_star"], E_rated=design.E_rated,
                capex_ess=cx, stack=fixed_stack, profile=profile)


def _size_ess(case, profile, fixed_stack, j, q_grid=None, verbose=True):
    is_c2 = case == "case2"
    name = C2_NAME if is_c2 else C3_NAME
    op = c2_operating_point(float(j), fixed_stack.A_tot)
    pstar = op["P_star"]
    prc = (c2_charge_rating if is_c2 else c3_charge_rating)(profile, pstar)
    grid = tuple(_q_grid(case) if q_grid is None else q_grid)
    if not grid or any(not 0 <= float(q) <= 100 for q in grid):
        raise ValueError("q_grid에는 0~100 사이의 분위수가 필요합니다")
    conf = CASE2 if is_c2 else CASE3
    usable = conf["soc_max"] - conf.get("soc_stop", conf["soc_floor"])
    thr = pstar / ETA_RTE if is_c2 else pstar
    need = daily_storage_need(profile.P_in, pstar, prc, profile.dt, threshold=thr)
    rows, best, evaluated = [], None, {}

    # case3 dispatch(dispatch_case3)는 ESS 용량이 P*의 극히 일부(3분 미만)로
    # 물리적으로 무의미할 만큼 작을 때, 시간별 이벤트 루프(range(100_000))가
    # 한 시간 안에서 수십만 번씩 완충/완방전을 반복하는 병리적 경우가 생겨
    # (2026-09 실측: 한 후보에서만 40초 이상, 전체 탐색이 100초 넘게 느려짐)
    # 계산이 크게 낭비된다. q-autoextend가 점점 더 작은 q를 시도하므로 이 하한이
    # 특히 중요하다. 이런 후보는 RTE 손실·재기동 비용 때문에 어차피 최적점이 될
    # 수 없으므로 dispatch를 돌리기 전에 걸러낸다 — 검증 결과 최적 조합·LCOH는
    # 완전히 동일하게 나오고 속도만 개선된다. case2는 이 문제가 관측되지 않아
    # 그대로 둔다(기존 1e-9 기준 유지).
    min_E_for_dispatch = (0.05 * pstar) if not is_c2 else 0.0

    def evaluate(q):
        nonlocal best
        E = float(np.percentile(need, q)) / usable
        if E <= max(1e-9, min_E_for_dispatch):
            return
        if E in evaluated:
            row = dict(evaluated[E], **{"q [%]": q})
            rows.append(row)
            return
        design = ESSDesign(name, float(j), E, prc, pstar, float(q))
        out = evaluate_fixed_design(profile, fixed_stack, design)
        r, cost = out["result"], out["lcoh"]["LCOH"]
        if r.H2_total <= 0 or not np.isfinite(cost):
            return
        row = {"q [%]": q, "j": float(j), "hours_of_P*": E / pstar, "E_rated [kWh]": E,
               "op_hours": _annual_mean(r, r.op_hours), "annual_H2 [t]": _annual_mean(r, r.H2_total) / 1000,
               "curtail [%]": r.E_curtail / r.E_paid * 100 if r.E_paid else 0,
               "RTE loss [%]": r.E_rte / r.E_paid * 100 if r.E_paid else 0,
               "restarts": _annual_mean(r, r.restarts), "eq_cycles": _annual_mean(r, r.eq_cycles),
               "CAPEX_ESS [kUSD]": out["capex_ess"] / 1e3,
               "LCOH_internal": cost, "LCOH_compare": cost,
               "P_star": pstar, "P_rated_chg": prc, "P_rated_dis": pstar}
        rows.append(row)
        evaluated[E] = row
        if best is None or cost < best[1]:
            best = (E / pstar, cost, E, r)

    for q in sorted(set(grid)):
        evaluate(q)
    if best is None:
        raise NoFeasibleDesignError(f"{name}: j={j:.4f}에서 유효한 ESS 후보가 없습니다")
    df = pd.DataFrame(rows).sort_values("q [%]").reset_index(drop=True)
    if not is_c2 and _ESS_AUTO_EXTEND_Q:
        for _ in range(3):
            qlo = float(df.iloc[0]["q [%]"])
            if len(df) <= 1 or int(df.LCOH_internal.idxmin()) != 0 or qlo <= 0.5:
                break
            for q in (qlo * 0.25, qlo * 0.5, qlo * 0.75):
                evaluate(round(q, 4))
            df = pd.DataFrame(rows).sort_values("q [%]").reset_index(drop=True)
    rippl = rippl_capacity(profile.P_in, pstar, ETA_RTE, profile.dt)
    if verbose:
        q = df.loc[df.LCOH_internal.idxmin(), "q [%]"]
        print(f"[{name} ESS] j={j:.4f}, E={best[2]:,.0f} kWh, q={q:g}%, LCOH={best[1]:.3f}")
    return best, df, rippl


def _search_multiyear_design(case, profile, fixed_stack, n_j=None, q_grid=None, verbose=True):
    jgrid = np.linspace(J_MIN, J_MAX, n_j or _n_j_grid(case))
    rows, best, best_df, best_rippl = [], None, None, None
    for i, j in enumerate(jgrid):
        try:
            candidate, df, rippl = _size_ess(case, profile, fixed_stack, j, q_grid, verbose=False)
        except NoFeasibleDesignError:
            continue
        rows.extend(df.to_dict("records"))
        if best is None or candidate[1] < best[1]:
            best, best_df, best_rippl = candidate, df, rippl
        if verbose and ((i + 1) % 4 == 0 or i + 1 == len(jgrid)):
            print(f"[{case} multi-year search] j {i+1}/{len(jgrid)}, best LCOH={best[1]:.3f}")
    if best is None:
        raise NoFeasibleDesignError(f"{case}: 유효한 공통 운전조건을 찾지 못했습니다")
    r = best[3]
    winner = best_df.loc[best_df.LCOH_internal.idxmin()]
    design = ESSDesign(r.case, float(r.j_star), r.E_rated, r.P_rated_chg, r.P_rated_dis,
                       float(winner["q [%]"]))
    op = c2_operating_point(design.j_star, fixed_stack.A_tot)
    detail = dict(j=design.j_star, key=design.sizing_q, hours=best[0], E_rated=best[2],
                  LCOH_internal=best[1], result=r, P_rc=r.P_rated_chg, P_rd=r.P_rated_dis,
                  op=op, design=design, mode="surplus_q")
    records = pd.DataFrame(rows)
    records["LCOH"] = records.LCOH_internal
    return dict(best=detail, records=records, j_star=design.j_star, cap_hours=best[0],
                cap_sweep=best_df, j_sweep=records.loc[records.groupby("j").LCOH_internal.idxmin()].reset_index(drop=True),
                rippl=best_rippl, mode="surplus_q", j_grid=jgrid, design=design)


def _run_ess_case(case, profile=None, fixed_stack=None, j_star=None, cap_hours=None,
                  optimize=True, verbose=True, *, design=None):
    profile = profile or get_profile()
    fixed_stack = fixed_stack or build_fixed_stack(profile, verbose=verbose)
    name = C2_NAME if case == "case2" else C3_NAME
    cap_df, opt, result = pd.DataFrame(), None, None
    rippl = float("nan")
    if design is not None:
        if design.case != name:
            raise ValueError("고정 설계의 Case가 실행 함수와 다릅니다")
        if j_star is not None or cap_hours is not None:
            raise ValueError("design과 j_star/cap_hours를 동시에 지정할 수 없습니다")
    else:
        if j_star is None and j_is_fixed():
            j_star = fixed_j(case)
        if j_star is not None and not J_MIN <= float(j_star) <= J_MAX:
            raise ValueError("j_star가 운전 창 밖입니다")
        if cap_hours is not None:
            if j_star is None or not np.isfinite(cap_hours) or cap_hours <= 0:
                raise ValueError("cap_hours 고정 평가에는 양수 cap_hours와 j_star가 필요합니다")
            op = c2_operating_point(j_star, fixed_stack.A_tot)
            prc = (c2_charge_rating if case == "case2" else c3_charge_rating)(profile, op["P_star"])
            design = ESSDesign(name, float(j_star), cap_hours * op["P_star"], prc, op["P_star"])
        elif not optimize:
            raise ValueError("optimize=False에는 design 또는 j_star+cap_hours가 필요합니다")
        elif j_star is None:
            opt = _search_multiyear_design(case, profile, fixed_stack, verbose=verbose)
            design, result, cap_df, rippl = opt["design"], opt["best"]["result"], opt["cap_sweep"], opt["rippl"]
        else:
            best, cap_df, rippl = _size_ess(case, profile, fixed_stack, j_star, verbose=verbose)
            result = best[3]
            winner = cap_df.loc[cap_df.LCOH_internal.idxmin()]
            design = ESSDesign(name, float(j_star), result.E_rated, result.P_rated_chg,
                               result.P_rated_dis, float(winner["q [%]"]))
    if result is None:
        out = evaluate_fixed_design(profile, fixed_stack, design)
    else:
        op = c2_operating_point(design.j_star, fixed_stack.A_tot)
        cx = ess_capex(design.E_rated, max(design.P_rated_chg, design.P_rated_dis))
        out = dict(result=result, lcoh=lcoh_compare(result, fixed_stack, cx), op=op,
                   design=design, j_star=design.j_star, cap_hours=design.E_rated / op["P_star"],
                   E_rated=design.E_rated, capex_ess=cx, stack=fixed_stack, profile=profile)
    out.update(cap_sweep=cap_df, rippl=rippl, opt=opt)
    if verbose:
        print(f"\n[{name}] {profile.period_label} common design: j*={design.j_star:.4f}, "
              f"E={design.E_rated:,.1f} kWh, Pchg/Pdis={design.P_rated_chg:,.1f}/{design.P_rated_dis:,.1f} kW")
        out["result"].report(h2_area_at_rated=float(h2_area(J_MAX)[0]),
                             h2_area_at_jstar=float(h2_area(design.j_star)[0]), A_tot=fixed_stack.A_tot)
        report_lcoh(name, out["lcoh"])
    return out



def size_ess_case2(profile, fixed_stack, j, q_grid=None, verbose=True):
    return _size_ess("case2", profile, fixed_stack, j, q_grid, verbose)


def find_operating_point_case2(profile, fixed_stack, n_round=2, verbose=True):
    return _search_multiyear_design("case2", profile, fixed_stack, verbose=verbose)



def run_case2(profile=None, fixed_stack=None, j_star=None, cap_hours=None,
              optimize=True, verbose=True, *, design=None):
    return _run_ess_case("case2", profile, fixed_stack, j_star, cap_hours,
                         optimize, verbose, design=design)



# %%
# PART 13 — CASE 3 : 하이브리드 (직결 + ESS)
########################################################################################
########################################################################################

C3_NAME = "Case 3"


def c3_operating_point(j: float, A_tot: float) -> dict:
    return c2_operating_point(j, A_tot)


def c3_charge_rating(profile, P_star: float) -> float:
    surplus = np.clip(profile.P_in - P_star, 0.0, None)
    if not _use_charge_pctl():
        return max(float(surplus.max()), 1e-6)
    pos = surplus[surplus > 0.0]
    if pos.size == 0:
        return 1e-6
    return max(float(np.percentile(pos, CHARGE_RATING_PCTL)), 1e-6)


def _c3_restart_level(E_rated, P_star, eta, dt, restart_hours,
                      soc_floor=None, soc_restart=None, soc_max=None):
    floor = CASE3["soc_floor"] if soc_floor is None else float(soc_floor)
    restart = CASE3["soc_restart"] if soc_restart is None else float(soc_restart)
    hi = CASE3["soc_max"] if soc_max is None else float(soc_max)
    if restart_hours is not None:
        return min(E_rated * floor + float(restart_hours) * P_star / eta, E_rated * hi)
    return E_rated * restart


def dispatch_case3(P_in, op, A_tot, E_rated, P_rated_chg, P_rated_dis,
                   soc_max=None, soc_floor=None, soc_stop=None, soc_restart=None,
                   restart_hours=None, eta_RTE=None, soc0=None, dt=1.0,
                   keep_series=True, running0=False) -> DispatchResult:
    hi = float(CASE3["soc_max"] if soc_max is None else soc_max)
    lo = float(CASE3["soc_floor"] if soc_floor is None else soc_floor)
    stop_soc = float(CASE3["soc_stop"] if soc_stop is None else soc_stop)
    restart_soc = float(CASE3["soc_restart"] if soc_restart is None else soc_restart)
    eta = float(ETA_RTE if eta_RTE is None else eta_RTE)
    P, E = _storage_inputs(P_in, op, E_rated, P_rated_chg, P_rated_dis, eta,
                           dt, hi, lo, soc0)
    E_hi, E_lo, E_stop = E_rated * hi, E_rated * lo, E_rated * stop_soc
    E_rs = _c3_restart_level(E_rated, op["P_star"], eta, dt, restart_hours, lo, restart_soc, hi)
    if not lo <= stop_soc < hi or (E_rated > 0 and not E_stop < E_rs <= E_hi):
        raise ValueError("Case 3는 floor <= stop < restart <= max의 히스테리시스가 필요합니다")
    Etol, Ttol = max(1e-10, E_rated * 1e-12), 1e-12
    Ptol = max(1e-10, op["P_star"] * 1e-12)
    running = bool(running0)
    tr = _StorageTrace(P, op, E_rated, E, eta, dt, keep_series, running)
    for i, p in enumerate(P):
        remaining = float(dt)
        deficit = max(op["P_star"] - float(p), 0.0)
        can_discharge = deficit <= P_rated_dis + Ptol
        direct_start = bool(CASE3.get("direct_restart", True)) and deficit <= Ptol
        for _ in range(100_000):
            if remaining <= Ttol:
                break
            if running and deficit > Ptol and (not can_discharge or E <= E_stop + Etol):
                running = False
                tr.stop(i, E)
            if not running:
                battery_start = (can_discharge and E >= E_rs - Etol and
                                 (deficit <= Ptol or E > E_stop + Etol))
                if direct_start or battery_start:
                    running = True
            direct = min(float(p), op["P_star"]) if running else 0.0
            d = deficit / eta if running else 0.0
            q = min(max(float(p) - direct, 0.0), P_rated_chg)
            if E >= E_hi - Etol:
                q = min(q, d)
            net = q - d
            h = remaining
            if running and net < 0:
                h = min(h, max(E - E_stop, 0.0) / (-net))
            if net > 0:
                h = min(h, max(E_hi - E, 0.0) / net)
                if not running and can_discharge and E < E_rs - Etol:
                    h = min(h, (E_rs - E) / net)
            if h <= Ttol:
                raise RuntimeError("Case 3 시간 적분이 진행하지 않습니다. SOC 조건을 확인하십시오")
            E += net * h
            remaining = max(remaining - h, 0.0)
            tr.record(i, h, running, q, d, direct, E)
            if E < E_lo - 10 * Etol or E > E_hi + 10 * Etol:
                raise RuntimeError("Case 3 SOC 수지 위반")
            if running and deficit > Ptol and E <= E_stop + Etol:
                running = False
                tr.stop(i, E)
        else:
            raise RuntimeError("Case 3 한 입력 구간의 이벤트 수가 과도합니다")
        tr.soc[i] = E / E_rated if E_rated > 0 else 0.0
    return tr.finish(C3_NAME, E, P_rated_chg, P_rated_dis, running,
                     {"E_hi": E_hi, "E_lo": E_lo, "E_stop": E_stop, "E_restart": E_rs,
                      "eta_RTE": eta,
                      "min_run_duration_h": (E_rs - E_stop) / (op["P_star"] / eta)})


def c3_run_warmstart(P_in, op, A_tot, E_rated, P_rc, P_rd, dt=1.0, **kw):
    return _run_storage_warmstart(dispatch_case3, (P_in, op, A_tot, E_rated, P_rc, P_rd),
                                  E_rated, dt, kw)


def find_operating_point_case3(profile, fixed_stack, n_j=None, q_grid=None, verbose=True):
    return _search_multiyear_design("case3", profile, fixed_stack, n_j, q_grid, verbose)



def size_ess_case3(profile, fixed_stack, j, q_grid=None, verbose=True):
    return _size_ess("case3", profile, fixed_stack, j, q_grid, verbose)


def run_case3(profile=None, fixed_stack=None, j_star=None, cap_hours=None,
              optimize=True, verbose=True, *, design=None):
    return _run_ess_case("case3", profile, fixed_stack, j_star, cap_hours,
                         optimize, verbose, design=design)



# %%
# PART 14 — 3케이스 비교 그래프
########################################################################################
########################################################################################

_CASE_ORDER = ["Case 1", "Case 2", "Case 3"]


def _production_hours(res, dt: float | None = None) -> dict:
    step = float(res.dt if dt is None else dt)
    p = np.asarray(res.P_in, dtype=float)
    on_h = (np.asarray(res.on_hours, dtype=float) if res.on_hours is not None
            else (np.asarray(res.j, dtype=float) > 1e-9).astype(float) * step)
    if on_h.shape != p.shape or np.any(on_h < -1e-9) or np.any(on_h > step + 1e-9):
        raise ValueError("on_hours는 입력 길이와 같고 각 구간의 [0, dt] 안이어야 합니다")
    on_h = np.clip(on_h, 0.0, step)
    stopped = step - on_h
    gen = p > 1e-9
    return {"producing": float(on_h.sum()),
            "idle_with_gen": float(stopped[gen].sum()),
            "no_gen": float(stopped[~gen].sum()),
            "mask": on_h > 1e-12, "on_hours": on_h,
            "generation_absent_hours": float((~gen).sum() * step)}


def _monthly(series, timestamps, agg="sum"):
    if agg != "sum":
        raise ValueError("월별 그래프는 연도-월 합계의 연평균을 사용합니다")
    stats = monthly_statistics(series, timestamps)
    return np.arange(1, 13), stats["mean"].to_numpy()


def _monthly_std(series, timestamps):
    return monthly_statistics(series, timestamps)["std"].to_numpy()


def _bar_top_labels(ax, x, values, fmt, fontsize=8):
    for xi, value in zip(x, values):
        ax.annotate(format(value, fmt), (xi, value), xytext=(0, 3),
                    textcoords="offset points", ha="center", va="bottom", fontsize=fontsize)



_RE_KIND_EN = {"태양광": "PV", "풍력": "Wind"}


def _profile_label(profile) -> str:
    raw = str(getattr(profile, "region", "") or "").strip()
    meta = getattr(profile, "meta", None) or {}
    kind = str(meta.get("kind", "") or "").strip()
    region = raw
    if kind and raw.endswith(kind):
        region = raw[: -len(kind)].strip()
    else:
        for k in _RE_KIND_EN:
            if raw.endswith(k):
                region, kind = raw[: -len(k)].strip(), k
                break
    return f"{region} {_RE_KIND_EN.get(kind, kind)}".strip()


def plot_g01_re_generation(profile, name="G01_RE_generation_monthly"):
    """연도별 월 발전량: 상자=25~75%, 중앙선=중앙값, 수염=최소~최대"""
    if plt is None:
        return
    table = monthly_year_table(profile.P_gen * profile.dt, profile.timestamps) / 1e3
    x = np.arange(1, 13)
    samples = [table[month].dropna().to_numpy(dtype=float) for month in x]
    if any(len(values) == 0 for values in samples):
        raise ValueError("발전량 박스플롯에는 1~12월의 연도별 월 합계가 필요합니다")

    fig, ax = plt.subplots(figsize=(8.0, 4.4))
    # 스타일 설정
    BOX_WIDTH = 0.40
    CAP_WIDTH = BOX_WIDTH * 0.50
    LINE_COLOR = "#000000"
    MID_COLOR = "#FFAE00"
    FILL_COLOR = "#FFAE0029"

    fig, ax = plt.subplots(figsize=(8.0, 4.4))
    boxes = ax.boxplot(
        samples,
        positions=x,
        widths=BOX_WIDTH,
        whis=(0, 100),
        patch_artist=True,
        showfliers=False,
        showmeans=False,
        manage_ticks=False,
        boxprops={
            "facecolor": FILL_COLOR,
            "edgecolor": LINE_COLOR,
            "linewidth": 0.8,
        },
        medianprops={
            "color": MID_COLOR, "linewidth": 1.8,
            "solid_capstyle": "butt",
        },
        whiskerprops={
            "color": LINE_COLOR, "linewidth": 1.3,
            "linestyle": "-",
        },
        capprops={
            "color": LINE_COLOR, "linewidth": 1.3,
            "solid_capstyle": "butt",
        },
        zorder=3,
    )
    
    for month, cap_low, cap_high in zip(
        x, boxes["caps"][::2], boxes["caps"][1::2]
    ):
        for cap in (cap_low, cap_high):
            cap.set_xdata([month - CAP_WIDTH / 2, month + CAP_WIDTH / 2])
    
    point_x = [month + (np.linspace(-0.12, 0.12, len(values))
                        if len(values) > 1 else np.zeros(1))
               for month, values in zip(x, samples)]
    ax.set_xticks(x)
    ax.set_xlim(0.4, 12.6)
    ax.set_xlabel("Month")
    ax.set_ylabel("Generation [MWh]")
    low = min(float(values.min()) for values in samples)
    high = max(float(values.max()) for values in samples)
    span = max(high - low, high * 0.10, 1.0)
    ax.set_ylim(max(0.0, low - 0.15 * span), high + 0.25 * span if high > 0 else 1.0)
    ax.xaxis.grid(False)
    ax.set_title(f"Monthly generation — {_profile_label(profile)} {profile.period_label}")
    ax.legend(
        [boxes["boxes"][0], boxes["whiskers"][0]],
        ["25–75% (line: median)", "Min–max"],
        loc="upper center", ncol=3, fontsize=8,
    )
    fig.tight_layout()
    _finish(fig, name)



def plot_case_operating_hours(case: str, res, months, idx: int):
    if plt is None:
        return
    months = res.timestamps if res.timestamps is not None else months
    nyr = _annual_n(res)
    h = _production_hours(res, getattr(res, "dt", 1.0) or 1.0)
    color = CASE_COLORS.get(case, "#333333")

    fig, axes = plt.subplots(1, 2, figsize=(11.0, 4.2))

    ax = axes[0]
    vals = [h[k] / nyr for k in ("producing", "idle_with_gen", "no_gen")]
    labels = ["producing H$_2$", "generation available,\nstopped", "no generation,\nstopped"]
    colors = [color, "#e9c46a", "#adb5bd"]
    w, _tx, _at = ax.pie(vals, labels=labels, colors=colors, startangle=90,
                         autopct=lambda p: f"{p:.1f}%\n({p / 100 * sum(vals):,.0f} h)",
                         pctdistance=0.72, textprops={"fontsize": 8},
                         wedgeprops={"width": 0.45, "edgecolor": "w"})
    ax.set_title(f"Annual hour ledger (total {sum(vals):,.0f} h)")

    ax = axes[1]
    x, mh = _monthly(h["on_hours"], months)
    ax.bar(x, mh, width=0.62, color=color)
    _bar_top_labels(ax, x, mh, ",.0f", 7)
    ax.set_xticks(x)
    ax.set_xlabel("Month")
    ax.set_ylabel("Producing hours [h]")
    ax.set_ylim(0, mh.max() * 1.18 if mh.max() > 0 else 1)
    ax.set_title("Monthly producing hours")

    fig.suptitle(f"{case} — hydrogen production hours "
                 f"({h['producing'] / nyr:,.0f} h/yr)", fontsize=12)
    fig.tight_layout()
    _finish(fig, f"G{idx:02d}_{case.replace(' ', '').lower()}_operating_hours")



def plot_case_monthly_h2(case: str, res, months, idx: int):
    if plt is None:
        return
    months = res.timestamps if res.timestamps is not None else months
    x, kg = _monthly(res.H2_series, months)
    t = kg / 1e3
    fig, ax = plt.subplots(figsize=(8.0, 4.2))
    ax.bar(x, t, width=0.62, color=CASE_COLORS.get(case, "#333333"))
    _bar_top_labels(ax, x, t, ".2f")
    ax.set_xticks(x)
    ax.set_xlabel("Month")
    ax.set_ylabel("H$_2$ [t]")
    ax.set_ylim(0, t.max() * 1.20 if t.max() > 0 else 1)
    ax.set_title(f"{case} — monthly H$_2$ (BOL) "
                 f"({_annual_mean(res, res.H2_total) / 1e3:,.2f} t/yr)")
    fig.tight_layout()
    _finish(fig, f"G{idx:02d}_{case.replace(' ', '').lower()}_monthly_H2")



def lcoh_vs_j_case1(profile, fixed_stack, n: int = 25) -> pd.DataFrame:
    j_lo, j_hi = max(2.0 * J_MIN, 0.4), 0.92 * J_LIM
    rows = []
    for jr in np.unique(np.r_[np.linspace(j_lo, j_hi, n), fixed_stack.j_rated]):
        st = size_fixed_stack(fixed_stack.P_nameplate, j_rated=float(jr))
        res = dispatch_case1(profile.P_in, st.A_tot, J_MIN, float(jr), dt=profile.dt, timestamps=profile.timestamps)
        lc = lcoh_compare(res, st, capex_ess=0.0)
        rows.append({"j": float(jr), "LCOH": lc["LCOH"], "H2_t": _annual_mean(res, res.H2_total) / 1e3,
                     "curtail_%": res.E_curtail / res.E_paid * 100 if res.E_paid else np.nan})
    return pd.DataFrame(rows)



def lcoh_vs_j_case23(profile, fixed_stack, cap_hours, case, n=17, *, design=None):
    rows = []
    grid = np.linspace(J_MIN, J_MAX, n)
    if design is not None:
        grid = np.unique(np.r_[grid, design.j_star])
    for j in grid:
        op = c2_operating_point(float(j), fixed_stack.A_tot)
        if design is None:
            E = float(cap_hours) * op["P_star"]
            prc = (c2_charge_rating if case == C2_NAME else c3_charge_rating)(profile, op["P_star"])
            d = ESSDesign(case, float(j), E, prc, op["P_star"])
        else:
            if case != design.case:
                raise ValueError("민감도 Case와 설계 Case가 다릅니다")
            if case == C2_NAME and op["P_star"] > design.P_rated_dis + 1e-8:
                continue
            d = ESSDesign(case, float(j), design.E_rated, design.P_rated_chg,
                          design.P_rated_dis, design.sizing_q)
        out = evaluate_fixed_design(profile, fixed_stack, d)
        r, lc = out["result"], out["lcoh"]
        if not np.isfinite(lc["LCOH"]):
            continue
        rows.append(dict(j=float(j), LCOH=lc["LCOH"], H2_t=_annual_mean(r, r.H2_total)/1e3,
                         E_rated=d.E_rated, P_rated_chg=d.P_rated_chg,
                         P_rated_dis=d.P_rated_dis, op_hours=_annual_mean(r, r.op_hours)))
    return pd.DataFrame(rows)



def plot_case_lcoh_vs_j(case: str, df: pd.DataFrame, j_used: float | None,
                        lcoh_used: float | None, idx: int, xlabel: str):
    if plt is None or df is None or not len(df):
        return
    color = CASE_COLORS.get(case, "#333333")
    fig, ax = plt.subplots(figsize=(8.0, 4.4))
    ax.plot(df["j"], df["LCOH"], color=color, lw=2.2, marker="o", ms=4,
            label="LCOH(j)")

    k = int(df["LCOH"].idxmin())
    j_opt, l_opt = df.loc[k, "j"], df.loc[k, "LCOH"]
    ax.axvline(j_opt, color=color, ls="--", lw=1.1, alpha=0.7)
    ax.plot([j_opt], [l_opt], marker="*", ms=15, color="#e76f51", zorder=5,
            label=f"minimum  j = {j_opt:.2f}, LCOH = {l_opt:.2f}")
    if j_used is not None and lcoh_used is not None:
        ax.plot([j_used], [lcoh_used], marker="D", ms=8, mfc="none",
                mec="#264653", mew=1.8, zorder=5,
                label=f"model used  j = {j_used:.2f}, LCOH = {lcoh_used:.2f}")

    ax.set_xlabel(xlabel)
    ax.set_ylabel("LCOH [USD/kg-H$_2$]")
    ax.set_title(f"{case} — LCOH vs current density")
    ax.legend(fontsize=8)
    fig.tight_layout()
    _finish(fig, f"G{idx:02d}_{case.replace(' ', '').lower()}_LCOH_vs_j")


def plot_g11_total_h2(results: dict, name="G11_compare_total_H2"):
    if plt is None:
        return
    cases = [case for case in _CASE_ORDER if results.get(case) is not None]
    if not cases:
        return
    tons = np.array([_annual_mean(results[case], results[case].H2_total) / 1e3 for case in cases], dtype=float)
    colors = [CASE_COLORS.get(case, "#333333") for case in cases]

    fig, ax = plt.subplots(figsize=(6.0, 4.8))
    bars = ax.bar(cases, tons, width=0.55, color=colors, edgecolor="white", zorder=3)

    base = float(tons[0]) if tons.size and tons[0] > 0 else float("nan")
    top = float(tons.max()) if tons.size else 1.0
    for b, v in zip(bars, tons):                      
        ax.text(b.get_x() + b.get_width() / 2, v, f"{v:,.2f} t",
                ha="center", va="bottom", fontsize=9, zorder=4)

    ax.set_ylabel("Annual H$_2$ production [t/yr]")
    ax.set_ylim(0, top * 1.24 if top > 0 else 1.0)
    ax.set_title("Annual hydrogen production (BOL)")
    ax.xaxis.grid(False)
    fig.tight_layout()
    _finish(fig, name)

    k = int(np.argmax(tons))
    print(f"\n[G10] Annual H2 production   (max : {cases[k]}, {tons[k]:,.2f} t/yr)")
    for i, (case, v) in enumerate(zip(cases, tons)):
        if i == 0 or not np.isfinite(base):
            print(f"  {case} : {v:,.2f} t/yr   (baseline)")
        else:
            print(f"  {case} : {v:,.2f} t/yr   "
                  f"({(v / base - 1) * 100:+.1f} % vs {cases[0]})")



def plot_g12_monthly_h2(results: dict, months, name="G12_compare_monthly_H2"):
    if plt is None:
        return
    fig, ax = plt.subplots(figsize=(9.0, 4.4))
    w = 0.26
    for i, case in enumerate(_CASE_ORDER):
        r = results.get(case)
        if r is None:
            continue
        x, kg = _monthly(r.H2_series, months)
        ax.bar(x + (i - 1) * w, kg / 1e3, width=w, color=CASE_COLORS.get(case),
               label=f"{case}  ({_annual_mean(r, r.H2_total) / 1e3:,.1f} t/yr)")
    ax.set_xticks(np.arange(1, 13))
    ax.set_xlabel("Month")
    ax.set_ylabel("H$_2$ [t]")
    ax.set_title("Monthly hydrogen production (BOL)")
    ax.legend(fontsize=9)
    fig.tight_layout()
    _finish(fig, name)



def plot_g13_operating_hours(results: dict, months, name="G13_compare_operating_hours"):
    if plt is None:
        return
    cases = [case for case in _CASE_ORDER if case in results]
    hrs = {case: _production_hours(results[case]) for case in cases}

    # --- (a) 연간 시간 구성 ------------------------------------------------------
    fig, ax = plt.subplots(figsize=(7.2, 4.2))
    segs = [("producing", "producing H$_2$", "#2a9d8f"),
            ("idle_with_gen", "generation available, stopped", "#e9c46a"),
            ("no_gen", "no generation, stopped", "#adb5bd")]
    lefts = np.zeros(len(cases))
    for key, label, color in segs:
        vals = np.array([hrs[case][key] / _annual_n(results[case]) for case in cases])
        ax.barh(cases, vals, left=lefts, color=color, label=label, height=0.5)
        for i, v in enumerate(vals):
            if v > 350:
                ax.text(lefts[i] + v / 2, i, f"{v:,.0f}", ha="center", va="center",
                        fontsize=9)
        lefts += vals
    ax.set_xlabel("Hours per year [h]")
    ax.set_xlim(0, lefts.max() * 1.02)
    ax.set_title("Annual hour breakdown")
    ax.legend(fontsize=8, loc="upper center", bbox_to_anchor=(0.5, -0.16), ncol=3)
    ax.invert_yaxis()
    fig.tight_layout()
    _finish(fig, f"{name}_a_annual")

    # --- (b) 월별 수소 생산 시간 -------------------------------------------------
    fig, ax = plt.subplots(figsize=(9.0, 4.2))
    w = 0.26
    for i, case in enumerate(cases):
        x, mh = _monthly(hrs[case]["on_hours"], months)
        ax.bar(x + (i - 1) * w, mh, width=w, color=CASE_COLORS.get(case),
               label=f"{case}  ({hrs[case]['producing'] / _annual_n(results[case]):,.0f} h/yr)")
    ax.set_xticks(np.arange(1, 13))
    ax.set_xlabel("Month")
    ax.set_ylabel("Producing hours [h]")
    ax.set_title("Monthly hydrogen producing hours")
    ax.legend(fontsize=8)
    fig.tight_layout()
    _finish(fig, f"{name}_b_monthly")



G14_LEDGER_LABELS = {"curtail": "curtailment", "idle": "below-minimum-load loss",
                     "rte": "ESS round-trip loss", "bop": "BOP",
                     "stack": "stack", "storage_delta": "ESS stored-energy change"}


def plot_g14_lcoh_and_ledger(results: dict, lcohs: dict, name="G14_LCOH_and_ledger", part="both"):
    if plt is None:
        return
    cases = [case for case in _CASE_ORDER if case in results]

    if part in ("both", "lcoh"):
        # --- (a) LCOH ---------------------------------------------------------------
        fig, ax = plt.subplots(figsize=(6.0, 4.8))
        bars = ax.bar(cases, [lcohs[case]["LCOH"] for case in cases],
                      color=[CASE_COLORS.get(case) for case in cases], width=0.55)
        for b, case in zip(bars, cases):
            ax.text(b.get_x() + b.get_width() / 2, b.get_height(), f"{b.get_height():.2f}",
                    ha="center", va="bottom", fontsize=10)
            share = lcohs[case].get("ess_share_pct", 0.0)
            if share and share > 0.05:
                ax.text(b.get_x() + b.get_width() / 2, b.get_height() * 0.5,
                        f"ESS\n{share:.1f}%", ha="center", va="center", fontsize=9, color="w")
        ax.set_ylabel("LCOH [USD/kg-H$_2$]")
        finite_costs = [lcohs[c]["LCOH"] for c in cases if np.isfinite(lcohs[c]["LCOH"])]
        ax.set_ylim(0, max(17.0, max(finite_costs, default=1.0) * 1.20))
        ax.set_title("LCOH (Levelized cost of hydrogen)")
        fig.tight_layout()
        _finish(fig, f"{name}_a_LCOH" if part == "both" else name)


    if part in ("both", "ledger"):
        # --- (b) 지불 에너지의 이동 경로 ---------------------------------------------
        fig, ax = plt.subplots(figsize=(6.0, 4.8))
        bottoms = np.zeros(len(cases))
        for k in ("stack", "bop", "rte", "idle", "curtail", "storage_delta"):
            vals = np.array([_annual_mean(results[case], results[case].ledger[k]) / 1e3 for case in cases])
            if np.max(np.abs(vals), initial=0.0) <= 1e-8:
                continue
            ax.bar(cases, vals, bottom=bottoms, color=LEDGER_COLORS[k],
                   label=G14_LEDGER_LABELS[k], width=0.55)
            for i, v in enumerate(vals):
                tot = _annual_mean(results[cases[i]], results[cases[i]].E_paid) / 1e3
                if tot > 0 and v / tot > 0.04:
                    ax.text(i, bottoms[i] + v / 2, f"{v / tot * 100:.1f}%",
                            ha="center", va="center", fontsize=8)
            bottoms += vals
        ax.set_ylabel("Paid energy [MWh/yr]")
        ax.set_ylim(0, bottoms.max() * 1.28)
        ax.set_title("Annual energy flow pathway")
        ax.legend(fontsize=8, loc="upper center", ncol=2)
        fig.tight_layout()
        _finish(fig, f"{name}_b_energy_pathway" if part == "both" else name)



_LCOH_ITEM_LABELS = {
    "capex": "CAPEX",
    "fixed_OM": "fixed O&M",
    "stack_replacement": "stack replacement",
    "battery_replacement": "battery replacement",
    "electricity": "electricity",
    "curtailment_penalty": "curtailment penalty",
    "variable_OM": "variable O&M",
    "water": "water",
}
_LCOH_ITEM_COLORS = {
    "capex": "#264653", "fixed_OM": "#2a9d8f",
    "stack_replacement": "#8ab17d", "battery_replacement": "#b07d62",
    "electricity": "#e9c46a", "curtailment_penalty": "#e76f51",
    "variable_OM": "#f4a261", "water": "#a8dadc",
}


def plot_g15_lcoh_breakdown(lcohs: dict, name="G15_LCOH_breakdown"):
    if plt is None:
        return
    cases = [case for case in _CASE_ORDER if case in lcohs]
    fig, ax = plt.subplots(figsize=(6.0, 4.8))
    bottoms = np.zeros(len(cases))
    for key, label in _LCOH_ITEM_LABELS.items():
        vals = np.array([lcohs[case]["items_usd_per_kg"].get(key, 0.0) for case in cases])
        if vals.sum() <= 1e-9:
            continue
        ax.bar(cases, vals, bottom=bottoms, width=0.55,
               color=_LCOH_ITEM_COLORS[key], label=label)
        for i, v in enumerate(vals):
            if v > 0.35:
                ax.text(i, bottoms[i] + v / 2, f"{v:.2f}", ha="center", va="center",
                        fontsize=8, color="w")
        bottoms += vals
    for i, case in enumerate(cases):
        ax.text(i, bottoms[i], f"{lcohs[case]['LCOH']:.2f}", ha="center", va="bottom",
                fontsize=11, fontweight="bold")
    ax.set_ylabel("LCOH [USD/kg-H$_2$]")
    ax.set_ylim(0, bottoms.max() * 1.30)
    ax.set_title("LCOH breakdown")
    ax.legend(fontsize=8, ncol=3, loc="upper center")
    fig.tight_layout()
    _finish(fig, name)


# ── 스택 열화 · 교체 lifecycle 그림 ────────────────────────────────

_OPEX_ITEMS = (
    ("electricity", "Electricity", "#e9c46a"),
    ("water", "Water", "#a8dadc"),
    ("fixed_OM", "Fixed O&M", "#2a9d8f"),
    ("variable_OM", "Variable O&M", "#f4a261"),
    ("curtailment_penalty", "Curtailment penalty", "#e76f51"),
    ("battery_replacement", "Battery replacement", "#b07d62"),
    ("stack_replacement", "Stack replacement", "#9467bd"),
)


def _lifecycle_df(case: str, lcohs: dict, results: dict, fixed_stack,
                  outs: dict | None = None):
    """케이스별 lifecycle 테이블."""
    lc = (lcohs or {}).get(case, {})
    life = lc.get("lifecycle")
    if life is not None and len(life):
        return life
    r = (results or {}).get(case)
    if r is None:
        return None
    out = (outs or {}).get(case) or {}
    try:
        return lifecycle_lcoh(r, fixed_stack, float(out.get("capex_ess", 0.0)))["lifecycle"]
    except Exception:
        return None


_DEG_CONTRIB = (("steady", "steady (j level)", "#2a9d8f"),
                ("ramp", "ramp (1h reference)", "#e9c46a"),
                ("start", "startup (extra)", "#d1495b"),
                ("stop", "shutdown", "#9c6644"))


def plot_g17_degradation_contrib(results: dict, name="G17_degradation_contribution"):
    if plt is None or not results:
        return
    cases = [c for c in _CASE_ORDER if c in results]
    if not cases:
        return
    degs = {c: ensure_degradation(results[c]) for c in cases}
    x = np.arange(len(cases))

    fig, ax = plt.subplots(figsize=(7.4, 5.0))
    totals = np.array([sum(degs[c]["annual_contrib_V"].values()) * 1e3 for c in cases])
    bottom = np.zeros(len(cases))
    for key, label, color in _DEG_CONTRIB:
        v = np.array([degs[c]["annual_contrib_V"][key] * 1e3 for c in cases])
        ax.bar(x, v, bottom=bottom, color=color, label=label, width=0.55)
        for xi in range(len(cases)):
            if v[xi] > 0.05 * max(totals[xi], 1e-12):
                ax.text(xi, bottom[xi] + v[xi] / 2,
                        f"{v[xi] / totals[xi] * 100:.0f}%",
                        ha="center", va="center", fontsize=8.5)
        bottom += v
    for xi, c in enumerate(cases):
        ax.text(xi, totals[xi] * 1.02,
                f"{totals[xi]:.1f} mV/yr\n"
                f"{degs[c]['rate_eff_uV_h']:.1f} $\\mu$V/h  |  "
                f"{degs[c]['stack_life_hours'] / 1e3:.1f} kh",
                ha="center", va="bottom", fontsize=8.5)
    ax.set_ylim(0, max(totals.max(), 1e-9) * 1.24)
    ax.set_xticks(x)
    ax.set_xticklabels(cases)
    ax.set_ylabel("Annual degradation [mV yr$^{-1}$]")
    ax.set_title("Annual degradation breakdown")
    ax.legend(fontsize=9)
    fig.tight_layout()
    _finish(fig, name)



def plot_case_lifecycle(case: str, life, idx: int, lc: dict | None = None):
    if plt is None or life is None or not len(life):
        return
    yr = life["year"].to_numpy(dtype=float)
    repl_years = yr[life["n_replacement"].to_numpy() > 0]

    fig, axes = plt.subplots(1, 2, figsize=(13.2, 4.7))

    # --- 좌 : 연간 OPEX + 교체 -------------------------------------------------
    ax = axes[0]
    bottom = np.zeros(len(life))
    for key, label, color in _OPEX_ITEMS:
        if key not in life.columns:
            continue
        v = life[key].to_numpy(dtype=float) / 1e6
        if np.nanmax(np.abs(v)) <= 1e-12:
            continue
        ax.bar(yr, v, bottom=bottom, color=color, label=label, width=0.72)
        bottom += v
    ax.set_xlabel("Project year")
    ax.set_ylabel("Annual cost [MUSD]")
    ax.set_xticks(yr.astype(int))
    ax.set_title(f"Annual OPEX and stack replacement — {case}")
    ax.legend(fontsize=7, ncol=2)

    # --- 우 : 생애 효율 + 수소 생산량 -----------------------------------------
    ax = axes[1]
    ax2 = ax.twinx()
    ax2.bar(yr, life["H2_kg"].to_numpy(dtype=float) / 1e3, width=0.72,
            color="#bcd4e6", label="H$_2$ production")
    ax2.set_ylabel("H$_2$ production [t yr$^{-1}$]")
    ax2.grid(False)
    ax.plot(yr, life["eff_HHV_pct"].to_numpy(dtype=float), "o-",
            color=CASE_COLORS.get(case, "#1f6f8b"), lw=2.0, ms=5,
            label="System efficiency (HHV, consumed)")
    ax.set_zorder(ax2.get_zorder() + 1)
    ax.patch.set_visible(False)
    for y in repl_years:
        ax.axvline(y, color="#9467bd", ls="--", lw=1.2, alpha=0.85)
    ax.set_xlabel("Project year")
    ax.set_ylabel("System efficiency (HHV) [%]")
    ax.set_xticks(yr.astype(int))
    h1, l1 = ax.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax.legend(h1 + h2, l1 + l2, fontsize=8, loc="lower left")
    ax.set_title("Lifecycle efficiency and H$_2$ production")

    fig.suptitle(f"Lifecycle cash flow — {case} (dashed = replacement year)",
                 fontsize=12)
    fig.tight_layout()
    _finish(fig, f"G{idx:02d}_{case.replace(' ', '').lower()}_lifecycle")


def _remove_obsolete_figures():
    """이전 버전에서 삭제·이름 변경된 그림만 정리합니다."""
    directory = _ensure_fig_dir()
    if directory is None:
        return
    obsolete = (
        "G10_compare_mean_annual_H2.png", "G11_mean_annual_energy_pathway.png",
        "G13_stack_replacement.png", "G14_LCOH_comparison.png",
        "G15_LCOH_breakdown.png", "G16_case1_lifecycle.png",
        "G17_case2_lifecycle.png", "G18_case3_lifecycle.png",
        "G19_case1_LCOH_vs_j.png", "G20_case2_LCOH_vs_j.png",
        "G21_case3_LCOH_vs_j.png",
    )
    for filename in obsolete:
        (Path(directory) / filename).unlink(missing_ok=True)


def make_all_figures(profile, results, lcohs, fixed_stack, out1, out2, out3, verbose=True):
    if plt is None:
        return {}
    setup_plots()
    _FIGURE_MANIFEST.clear()
    _remove_obsolete_figures()
    ts = profile.timestamps
    outs = {"Case 1": out1, "Case 2": out2, "Case 3": out3}
    plot_g01_re_generation(profile, "G01_RE_generation_monthly")
    for i, case in enumerate(_CASE_ORDER):
        if case in results:
            plot_case_operating_hours(case, results[case], ts, 2 + i)
    plot_g13_operating_hours(results, ts, "G05_compare_operating_hours")
    for i, case in enumerate(_CASE_ORDER):
        if case in results:
            plot_case_monthly_h2(case, results[case], ts, 6 + i)
    plot_g12_monthly_h2(results, ts, "G09_compare_monthly_H2")
    plot_g11_total_h2(results, "G10_compare_annual_H2")
    plot_g14_lcoh_and_ledger(results, lcohs, "G11_annual_energy_pathway", part="ledger")
    plot_g17_degradation_contrib(results, "G12_degradation_contribution")
    plot_g14_lcoh_and_ledger(results, lcohs, "G13_LCOH_comparison", part="lcoh")
    plot_g15_lcoh_breakdown(lcohs, "G14_LCOH_breakdown")
    for i, case in enumerate(_CASE_ORDER):
        if case in results:
            plot_case_lifecycle(case, _lifecycle_df(case, lcohs, results, fixed_stack, outs),
                                15 + i, lcohs.get(case))
    sweeps = {}
    if verbose:
        print("  LCOH 민감도 계산 (G18~G20)")
    if "Case 1" in results:
        df = lcoh_vs_j_case1(profile, fixed_stack)
        sweeps["Case 1"] = df
        plot_case_lcoh_vs_j("Case 1", df, J_MAX, lcohs["Case 1"]["LCOH"], 18,
                            "Rated current density [A/cm$^2$] (stack resized)")
    for i, case in enumerate(("Case 2", "Case 3")):
        if case not in results:
            continue
        out = outs[case]
        df = lcoh_vs_j_case23(profile, fixed_stack, out["cap_hours"], case, design=out["design"])
        sweeps[case] = df
        plot_case_lcoh_vs_j(case, df, out["j_star"], lcohs[case]["LCOH"], 19 + i,
                            "Operating current density [A/cm$^2$] (PEM and ESS fixed)")
    return sweeps



def _ess_sizing_key(out: dict | None) -> str:
    if not out:
        return "-"
    df = out.get("cap_sweep")
    if isinstance(df, pd.DataFrame) and len(df):
        k = int(df["LCOH_internal"].idxmin())
        return f"q = {df.loc[k, 'q [%]']:g} % (일별 잉여 분위수)"
    design = out.get("design")
    if design is not None and design.sizing_q is not None:
        return f"q = {design.sizing_q:g} % (다년 일별 잉여 분위수)"
    return "고정 설비 평가"


def ess_summary_table(out2: dict | None, out3: dict | None) -> pd.DataFrame:
    rows = [{"case": "Case 1",
             "ESS 용량 E_rated [kWh]": 0.0,
             "ESS 용량 [MWh]": 0.0,
             "등가 저장시간 [h of P*]": 0.0,
             "ESS 출력 P_rated [kW]": 0.0,
             "ESS CAPEX [USD]": 0.0,
             "등가 사이클 [cyc/yr]": 0.0,
             "운전 j* [A/cm2]": float("nan"),
             "사이징 근거": "ESS 없음 (직결)"}]

    for case, out in (("Case 2", out2), ("Case 3", out3)):
        if not out:
            continue
        r = out.get("result")
        E = float(out.get("E_rated", getattr(r, "E_rated", 0.0)) or 0.0)
        P_star = float((out.get("op") or {}).get("P_star", 0.0) or 0.0)
        if P_star <= 0:
            P_star = float(getattr(r, "P_star", 0.0) or 0.0)
        P_rated = max(float(getattr(r, "P_rated_chg", 0.0) or 0.0),
                      float(getattr(r, "P_rated_dis", 0.0) or 0.0))
        rows.append({
            "case": case,
            "ESS 용량 E_rated [kWh]": E,
            "ESS 용량 [MWh]": E / 1e3,
            "등가 저장시간 [h of P*]": E / P_star if P_star > 0 else float("nan"),
            "ESS 출력 P_rated [kW]": P_rated,
            "ESS CAPEX [USD]": float(out.get("capex_ess", 0.0) or 0.0),
            "등가 사이클 [cyc/yr]": _annual_mean(r, float(getattr(r, "eq_cycles", 0.0) or 0.0)),
            "운전 j* [A/cm2]": float(out.get("j_star", float("nan"))),
            "사이징 근거": _ess_sizing_key(out),
        })
    return pd.DataFrame(rows).set_index("case")


def print_ess_summary(out2: dict | None, out3: dict | None,
                      title: str = "ESS 용량 최적화 결과 (E_rated*)",
                      save_csv: bool | str = False) -> pd.DataFrame:
    df = ess_summary_table(out2, out3)
    _banner(title, "-")
    print(f"  mode=surplus_q, c_E={C_E_USD_KWH} $/kWh, "
          f"c_P={C_P_USD_KW} $/kW, RTE={ETA_RTE:.2f}")
    cols = ["ESS 용량 E_rated [kWh]", "ESS 출력 P_rated [kW]",
            "ESS CAPEX [USD]", "등가 사이클 [cyc/yr]", "운전 j* [A/cm2]", "사이징 근거"]
    print(df[cols].to_string(float_format=lambda x: f"{x:,.2f}"))
    e2 = df.loc["Case 2", "ESS 용량 E_rated [kWh]"] if "Case 2" in df.index else np.nan
    e3 = df.loc["Case 3", "ESS 용량 E_rated [kWh]"] if "Case 3" in df.index else np.nan
    if np.isfinite(e2) and np.isfinite(e3) and e2 > 0:
        print(f"  Case3/Case2 E_rated={e3 / e2 * 100:,.1f} % "
              f"(reduction={100 - e3 / e2 * 100:,.1f} %)")
    if save_csv:
        try:
            if isinstance(save_csv, str):
                path = save_csv
            else:
                d = _ensure_fig_dir()
                path = f"{d}/ESS_sizing_summary.csv" if d else "ESS_sizing_summary.csv"
            df.to_csv(path, encoding="utf-8-sig")
            print(f"  saved: {path}")
        except Exception as e:
            print(f"  [WARN] ESS summary CSV save failed: {e}")
    return df


def lifecycle_summary_table(lcohs: dict) -> pd.DataFrame:
    rows = []
    for case in _CASE_ORDER:
        lc = (lcohs or {}).get(case)
        if not lc:
            continue
        life = lc.get("lifecycle")
        rows.append({
            "case": case,
            "LCOH [$/kg]": lc.get("LCOH", np.nan),
            "BOL 연평균 H2 [t]": lc.get("annual_H2", np.nan) / 1e3,
            "생애 H2 합계 [t]": lc.get("H2_total_life_kg", np.nan) / 1e3,
            "생애 평균 H2 [t/yr]": lc.get("annual_H2_mean", np.nan) / 1e3,
            "최종연도 열화손실 [%]": lc.get("H2_fade_pct", np.nan),
            "스택 교체 [회]": lc.get("replacement_count", np.nan),
            "평균 유지율 [%]": lc.get("retention_mean_pct", np.nan),
            "효율 yr1 [%HHV]": (float(life["eff_HHV_pct"].iloc[0])
                               if life is not None and len(life) else np.nan),
            "효율 최종 [%HHV]": (float(life["eff_HHV_pct"].iloc[-1])
                                if life is not None and len(life) else np.nan),
            "교체비 현재가치 [k$]": lc.get("items_usd", {}).get("stack_replacement", np.nan) / 1e3,
        })
    return pd.DataFrame(rows).set_index("case") if rows else pd.DataFrame()


def print_lifecycle_summary(lcohs: dict, title: str = "스택 열화 · 교체 lifecycle 결과") -> pd.DataFrame:
    df = lifecycle_summary_table(lcohs)
    _banner(title, "-")
    if not len(df):
        print("  (lifecycle 결과 없음)")
        return df
    print(f"  kernel=Su et al. 2024 (anchored), applies_to=stack, "
          f"앵커={STACK_LIFETIME_H:,.0f} h @ EOL 유지율 {STACK_EOL_RETENTION * 100:.0f} % "
          f"({deg_anchor_rate_uV_h():.2f} uV/h @ j={DEG_J_REF:.2f}), "
          f"프로젝트={LIFETIME_YR} yr, 교체 정지={REPLACEMENT_DOWNTIME_H:,.0f} h/회")
    print(df.to_string(float_format=lambda x: f"{x:,.2f}"))
    return df



def _require_summary_excel_engine() -> None:
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:
        raise ImportError(
            "최종 요약 Excel 저장에는 openpyxl이 필요합니다. "
            "실행 환경에서 'python -m pip install openpyxl'을 실행하세요. "
            "Excel 저장을 끄려면 SAVE_FINAL_SUMMARY_XLSX = False로 설정하세요."
        ) from exc


def _summary_filename_part(value, fallback: str) -> str:
    import re
    name = _normalize_re_name(value)
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name)
    name = re.sub(r"\s+", "_", name).strip(" ._")
    # UTF-8 파일명 길이도 고려해 지역/발전원 각각 최대 40자로 제한합니다.
    return name[:40] or fallback


def _save_new_summary_file(directory: Path, stem: str, content: bytes) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    number = 0
    while True:
        suffix = "" if number == 0 else f"_{number:03d}"
        path = directory / f"{stem}{suffix}.xlsx"
        try:
            stream = path.open("xb")
        except FileExistsError:
            number += 1
            continue
        try:
            with stream:
                stream.write(content)
        except BaseException:
            # 이번 호출이 새로 만든 불완전 파일만 제거합니다. 이전 파일은 건드리지 않음.
            path.unlink(missing_ok=True)
            raise
        return path


def export_final_summary_excel(profile, fixed_stack, metrics: pd.DataFrame,
                               lcoh_items: pd.DataFrame,
                               winner: tuple, best_h2: tuple) -> Path:
    from datetime import datetime
    from io import BytesIO

    _require_summary_excel_engine()
    from openpyxl.styles import Alignment, Border, Font, PatternFill
    from openpyxl.utils import get_column_letter

    kind = _normalize_re_name(profile.meta.get("kind") or RE_KIND)
    region = _normalize_re_name(profile.region or RE_REGION_NAME)
    if region.endswith(" " + kind):
        region = region[:-(len(kind) + 1)].strip()
    if not region or region == "(unset)":
        region = _normalize_re_name(RE_REGION_NAME)
    stem = (f"{_summary_filename_part(region, 'region')}_"
            f"{_summary_filename_part(kind, 'renewable')}")

    directory = Path(FINAL_SUMMARY_DIR).expanduser()
    if not directory.is_absolute():
        base = Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
        directory = base / directory

    config_rows = [
        ("저장 시각", datetime.now().astimezone().isoformat(timespec="seconds"), "실행 환경 현지시각 / UTC offset 포함"),
        ("지역", region, "실제 입력 프로파일"),
        ("재생에너지 종류", kind, "실제 입력 프로파일"),
        ("입력 표시명", profile.region, "input"),
        ("분석 기간", profile.period_label, "관측연도"),
        ("분석 연수", profile.n_years, "yr"),
        ("입력 출처", profile.source, "source"),
        ("재생에너지 용량", profile.meta.get("re_capacity_MW", np.nan), "MW; 미지정은 빈칸"),
        ("연평균 배정 전력량", profile.E_paid / profile.n_years / 1e3, "MWh/yr"),
        ("이용률 CF", profile.meta.get("unit_CF", np.nan) * 100, "% ; 최종 요약과 동일 기준"),
        ("배정 비율", profile.allocation_n * 100, "%"),
        ("시간 간격", profile.dt, "h"),
        ("PEM 명판 용량", fixed_stack.P_nameplate / 1000, "MW"),
        ("전류밀도 하한 J_MIN", J_MIN, "A/cm2"),
        ("전류밀도 상한 J_MAX", J_MAX, "A/cm2"),
        ("전기 단가 ELEC_PRICE", ELEC_PRICE, "$/kWh"),
        ("LCOH 기준", "lifecycle", "최종 요약 기준"),
        ("열화 모델", "Su2024/anchored", "stack"),
        ("EOL 유지율", STACK_EOL_RETENTION * 100, "%"),
        ("스택 수명 앵커", STACK_LIFETIME_H, "h"),
        ("스택 교체비 비율", REPL_STACK_FRAC * 100, "% of capex_stack"),
        ("프로젝트 기간", LIFETIME_YR, "yr"),
        ("할인율", INTEREST_RATE * 100, "%"),
        ("전류밀도 설정 모드", J_STAR_MODE, "J_STAR_MODE"),
        ("빠른 계산 모드", FAST_MODE, "FAST_MODE"),
        ("저장 제외 범위", "별도 ESS sizing 표", "case metrics의 ESS 지표는 유지"),
        ("빈칸 의미", "해당 없음 / 미계산 / 미지정", "원래 NaN인 셀은 0으로 바꾸지 않음"),
    ]
    config = pd.DataFrame(config_rows, columns=["항목", "값", "단위 / 설명"])
    conclusions = pd.DataFrame([
        {"항목": "LCOH 최저", "case": winner[0], "값": winner[1], "단위": "$/kg"},
        {"항목": "생산량 최대", "case": best_h2[0], "값": best_h2[1] / 1e3, "단위": "t/yr (BOL)"},
    ])

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        config.to_excel(writer, sheet_name="설정값", index=False)
        metrics.to_excel(writer, sheet_name="case_metrics", index=True, index_label="항목")
        lcoh_items.to_excel(writer, sheet_name="LCOH_breakdown", index=True, index_label="항목 [$/kg]")
        conclusions.to_excel(writer, sheet_name="결론", index=False)

        header_fill = PatternFill("solid", fgColor="17365D")
        total_fill = PatternFill("solid", fgColor="E2EFDA")
        for sheet in writer.sheets.values():
            sheet.sheet_view.showGridLines = False
            sheet.freeze_panes = "B2" if sheet.title in ("case_metrics", "LCOH_breakdown") else "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.row_dimensions[1].height = 26
            for row in sheet.iter_rows():
                for cell in row:
                    cell.font = Font(name="Calibri", size=11)
                    cell.border = Border()
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    if cell.data_type == "f":
                        cell.data_type = "s" 
                    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        cell.number_format = ("#,##0.0000" if sheet.title in ("설정값", "LCOH_breakdown")
                                              else "#,##0.000" if sheet.title == "결론" else "#,##0.00")
                if row[0].row > 1:
                    sheet.row_dimensions[row[0].row].height = 30 if sheet.title == "설정값" else 24
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            widths = ([38, 46, 52] if sheet.title == "설정값" else
                      [24, 18, 22, 22] if sheet.title == "결론" else
                      [44] + [20] * (sheet.max_column - 1))
            for i, width in enumerate(widths, start=1):
                sheet.column_dimensions[get_column_letter(i)].width = width
            if sheet.title == "LCOH_breakdown":
                for cell in sheet[sheet.max_row]:
                    cell.fill = total_fill
                    cell.font = Font(name="Calibri", size=11, bold=True)

    return _save_new_summary_file(directory, stem, buffer.getvalue())


def print_final_summary(profile, fixed_stack, results: dict, lcohs: dict,
                        out1: dict, out2: dict, out3: dict, sweeps: dict | None = None):
    _banner("최종 요약")
    cap_mw = profile.meta.get("re_capacity_MW", np.nan)
    cf_txt = (f", CF={profile.meta['unit_CF'] * 100:,.2f} %"
              if "unit_CF" in profile.meta else "")
    print(f"  input={profile.region} {profile.period_label}, "
          f"RE={cap_mw:,.2f} MW" if np.isfinite(cap_mw) else
          f"  input={profile.region} {profile.period_label}, RE=(unspecified)")
    print(f"  mean annual E_paid={profile.E_paid / profile.n_years / 1e3:,.1f} MWh{cf_txt}, "
          f"PEM={fixed_stack.P_nameplate / 1000:,.3f} MW, "
          f"j_window={J_MIN:.3f}-{J_MAX:.2f} A/cm2, "
          f"elec={ELEC_PRICE:,.4f} $/kWh")
    print(f"  LCOH 기준=lifecycle | 열화=Su2024/anchored"
          f"(EOL {STACK_EOL_RETENTION * 100:.0f} % @ {STACK_LIFETIME_H:,.0f} h), "
          f"교체비={REPL_STACK_FRAC * 100:.0f} % of capex_stack")

    rows = []
    for case, out in [("Case 1", out1), ("Case 2", out2), ("Case 3", out3)]:
        if case not in results:
            continue
        r = results[case]
        h = _production_hours(r, r.dt or 1.0)
        lc = lcohs[case]
        j_star = out.get("j_star")
        sw = (sweeps or {}).get(case)
        j_opt = float(sw.loc[sw["LCOH"].idxmin(), "j"]) if sw is not None and len(sw) else np.nan
        rows.append({
            "case": case,
            "연평균 생산시간 [h/yr]": _annual_mean(r, h["producing"]),
            "연평균 정지(발전有) [h/yr]": _annual_mean(r, h["idle_with_gen"]),
            "연평균 H2 [t] (BOL)": _annual_mean(r, r.H2_total) / 1e3,
            "생애 H2 [t]": lc.get("H2_total_life_kg", np.nan) / 1e3,
            "스택 교체 [회]": lc.get("replacement_count", np.nan),
            "평균 생산율 [kg/h]": r.H2_total / h["producing"] if h["producing"] else np.nan,
            "운전 j* [A/cm2]": j_star if j_star is not None else np.nan,
            "민감도 최소 j [A/cm2]": j_opt,
            "ESS 용량 [kWh]": r.E_rated,
            "ESS 출력 [kW]": max(r.P_rated_chg, r.P_rated_dis),
            "ESS CAPEX [k$]": lc["capex_ess"] / 1e3,
            "SEC_eff [kWh/kg]": r.SEC_eff(),
            "curtail [%]": r.E_curtail / max(r.E_paid, 1e-9) * 100,
            "LCOH [$/kg]": lc["LCOH"],
        })
    df = pd.DataFrame(rows).set_index("case").T
    print("\n[case metrics]")
    print(df.to_string(float_format=lambda x: f"{x:,.2f}"))

    items = pd.DataFrame({case: lcohs[case]["items_usd_per_kg"] for case in results}).round(4)
    items.index = [_LCOH_ITEM_LABELS.get(i, i) for i in items.index]
    items.loc["TOTAL"] = [lcohs[case]["LCOH"] for case in results]
    print(f"\n[LCOH breakdown, $/kg]  (기준: lifecycle)")
    print(items.to_string(float_format=lambda x: f"{x:,.3f}"))

    ess = ess_summary_table(out2, out3)
    ecols = ["ESS 용량 E_rated [kWh]", "ESS 출력 P_rated [kW]",
             "ESS CAPEX [USD]", "등가 사이클 [cyc/yr]", "사이징 근거"]
    print(f"\n[ESS sizing] mode=surplus_q")
    print(ess[ecols].to_string(float_format=lambda x: f"{x:,.2f}"))

    win = min(((case, lcohs[case]["LCOH"]) for case in results), key=lambda kv: kv[1])
    best_h2 = max(((case, _annual_mean(results[case], results[case].H2_total)) for case in results), key=lambda kv: kv[1])
    print(f"\n[결론] LCOH 최저={win[0]} ({win[1]:,.3f} $/kg), "
          f"생산량 최대={best_h2[0]} ({best_h2[1] / 1e3:,.2f} t/yr)")

    df.attrs["excel_path"] = None
    if SAVE_FINAL_SUMMARY_XLSX:
        excel_path = export_final_summary_excel(profile, fixed_stack, df, items, win, best_h2)
        df.attrs["excel_path"] = str(excel_path)
        print(f"\n[최종 요약 Excel 저장] {excel_path}")
    return df


def _banner(title: str, ch: str = "=") -> None:
    print("\n" + ch * 78)
    print(f" {title}")
    print(ch * 78)


def echo_config(profile, fixed_stack, win) -> None:
    _banner("공유 파라미터 요약 (PART 1)", "-")
    print(f"  catalyst=IrO2, T={T_OPER - 273.15:.0f} C, "
          f"membrane={MEMB_T_UM:.0f} um, BOP={BOP_EFFICIENCY:.2f} kWh/kg")
    print(f"  stack={fixed_stack.P_nameplate:,.1f} kW, "
          f"A_tot={fixed_stack.A_tot:,.0f} cm2, j={J_MIN:.4f}-{J_MAX:.2f} A/cm2, "
          f"P={win.P_min:,.0f}-{win.P_max:,.0f} kW")
    print(f"  RE={profile.region} {profile.period_label}, mean annual E_paid={profile.E_paid / profile.n_years / 1e3:,.1f} MWh, "
          f"n={profile.allocation_n * 100:.0f} %, elec={ELEC_PRICE} $/kWh, "
          f"ESS(c_E/c_P/RTE)={C_E_USD_KWH}/{C_P_USD_KW}/{ETA_RTE:.2f}")
    print(f"  LCOH=lifecycle, 열화=Su2024/anchored/stack "
          f"(EOL {STACK_EOL_RETENTION * 100:.0f} % @ {STACK_LIFETIME_H:,.0f} h), "
          f"프로젝트={LIFETIME_YR} yr, i={INTEREST_RATE * 100:.1f} %")


def setup_common(verbose: bool = True):
    global _OUTPUT_RUN_TAG
    profile = get_profile()
    _OUTPUT_RUN_TAG = f"{profile.region.replace(' ', '_')}_{profile.years[0]}_{profile.years[-1]}"
    if verbose:
        _banner("공통 셋업")
        profile.report()
    fixed_stack = build_fixed_stack(profile, verbose=verbose)
    win = operating_window(fixed_stack.A_tot)
    if verbose:
        echo_config(profile, fixed_stack, win)
    return profile, fixed_stack, win



def export_multiyear_results(profile, results, lcohs, outs, comparison=None, sweeps=None):
    directory = _ensure_fig_dir()
    if directory is None or not SAVE_RESULT_CSV:
        return {}
    d = Path(directory)
    saved = {}

    def save(name, table, index=True):
        path = d / name
        table.to_csv(path, index=index, encoding="utf-8-sig")
        saved[name] = str(path)

    for case, r in results.items():
        key = case.replace(" ", "").lower()
        annual = annual_result_table(r)
        save(f"{key}_annual_BOL.csv", annual)
        save(f"{key}_monthly_H2_by_year.csv", monthly_year_table(r.H2_series, profile.timestamps))
        save(f"{key}_monthly_H2_statistics.csv", monthly_statistics(r.H2_series, profile.timestamps))
        save(f"{key}_monthly_hours_by_year.csv", monthly_year_table(r.on_hours, profile.timestamps))
        save(f"{key}_monthly_hours_statistics.csv", monthly_statistics(r.on_hours, profile.timestamps))
        save(f"{key}_lifecycle.csv", lcohs[case]["lifecycle"], index=False)
        hourly = pd.DataFrame({"timestamp": profile.timestamps, "P_in_kW": r.P_in,
                               "j_A_cm2": r.j, "P_used_kW": r.P_used,
                               "H2_kg": r.H2_series, "producing_hours": r.on_hours})
        if r.soc is not None:
            hourly["SOC"] = r.soc
        save(f"{key}_hourly_BOL.csv", hourly, index=False)
        opt = outs[case].get("opt")
        if opt is not None and isinstance(opt.get("records"), pd.DataFrame):
            save(f"{key}_design_search.csv", opt["records"], index=False)
        if sweeps and case in sweeps:
            save(f"{key}_LCOH_vs_j.csv", sweeps[case], index=False)
    generation = profile.P_gen * profile.dt
    save("RE_annual_generation_kWh.csv", monthly_year_table(generation, profile.timestamps).sum(axis=1).to_frame("generation_kWh"))
    save("RE_monthly_generation_by_year_kWh.csv", monthly_year_table(generation, profile.timestamps))
    save("RE_monthly_generation_statistics_kWh.csv", monthly_statistics(generation, profile.timestamps))
    save("input_quality.csv", pd.DataFrame(profile.meta.get("quality", [])), index=False)
    save("comparison.csv", comparison if comparison is not None else compare_table(results, lcohs), index=False)
    if _FIGURE_MANIFEST:
        save("figure_order.csv", pd.DataFrame(_FIGURE_MANIFEST), index=False)
    import json
    manifest = dict(region=profile.region, source_years=list(profile.years),
        project_years=LIFETIME_YR, project_source_years=lcohs[next(iter(lcohs))]["lifecycle"].source_year.astype(int).tolist(),
        annual_statistics="calendar-year means; sample SD ddof=1", monthly_basis="BOL; no degradation derating",
        lifecycle_basis="repeat source-year BOL totals; annual-rate post-processing; not degradation-coupled hourly dispatch",
        warm_start_days=WARM_START_DAYS, warm_start_scope="once before full observed sequence; excluded",
        initial_storage={c: r.E_start for c, r in results.items()},
        final_storage={c: r.E_end for c, r in results.items()},
        source_paths=[str(Path(p).relative_to(Path(__file__).resolve().parent))
                      if Path(p).is_relative_to(Path(__file__).resolve().parent) else str(p)
                      for p in profile.meta.get("paths", [])], dt_hours=profile.dt,
        input_validation=dict(missing_policy=RE_MISSING_POLICY,
            low_cf_policy=RE_LOW_CF_POLICY, min_annual_cf=RE_MIN_ANNUAL_CF,
            low_cf_years=profile.meta.get("low_cf_years", []),
            zero_generation_years=profile.meta.get("zero_generation_years", [])),
        physical_parameters=dict(T_OPER=T_OPER, J_MIN=J_MIN, J_MAX=J_MAX,
            STACK_NAMEPLATE_KW=STACK_NAMEPLATE_KW, RE_CAPACITY_MW=RE_CAPACITY_MW,
            ALLOCATION_N=ALLOCATION_N, CASE2=CASE2, CASE3=CASE3),
        economic_parameters=dict(ELEC_PRICE=ELEC_PRICE, INTEREST_RATE=INTEREST_RATE,
            C_E_USD_KWH=C_E_USD_KWH, C_P_USD_KW=C_P_USD_KW),
        designs={c: dict(j_star=r.j_star, E_rated=r.E_rated, P_rated_chg=r.P_rated_chg,
                        P_rated_dis=r.P_rated_dis) for c, r in results.items()})
    path = d / "run_config.json"
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    saved[path.name] = str(path)
    return saved


def main(make_plots: bool = True, run_case1_sweep: bool = True) -> dict:
    if SAVE_FINAL_SUMMARY_XLSX:
        _require_summary_excel_engine()
    profile, fixed_stack, win = setup_common()

    _banner("케이스별 독립 실행 — 각 케이스를 '자기 최적 운전점'에서 돌린다")
    out1 = run_case1(profile, fixed_stack)
    rating_df = (stack_rating_sweep_case1(profile, base_nameplate=fixed_stack.P_nameplate)
                 if run_case1_sweep else None)
    out2 = run_case2(profile, fixed_stack)
    out3 = run_case3(profile, fixed_stack)

    _banner("케이스별 운전점")
    print(pd.DataFrame([
        {"case": "Case 1", "j*": "N/A (전력 추종)", "E_rated [kWh]": 0.0,
         "LCOH [$/kg]": out1["lcoh"]["LCOH"]},
        {"case": "Case 2", "j*": f"{out2['j_star']:.4f}",
         "E_rated [kWh]": out2["E_rated"], "LCOH [$/kg]": out2["lcoh"]["LCOH"]},
        {"case": "Case 3", "j*": f"{out3['j_star']:.4f}",
         "E_rated [kWh]": out3["E_rated"], "LCOH [$/kg]": out3["lcoh"]["LCOH"]},
    ]).to_string(index=False, float_format=lambda x: f"{x:,.3f}"))

    j_headline = {"Case 2": float(out2["j_star"]), "Case 3": float(out3["j_star"])}
    ess_df = print_ess_summary(out2, out3,
                               title="ESS 용량 최적화 결과 — 헤드라인",
                               save_csv=True)

    results = {"Case 1": out1["result"], "Case 2": out2["result"],
               "Case 3": out3["result"]}
    lcohs = {"Case 1": out1["lcoh"], "Case 2": out2["lcoh"], "Case 3": out3["lcoh"]}

    _banner("PEM 열화 · 스택 수명  (Su et al. 2024 커널)")
    degradation_df = print_degradation_summary(results)

    life_df = print_lifecycle_summary(lcohs)

    _banner("결론 비교표")
    cmp_df = compare_table(results, lcohs)
    print(cmp_df.to_string(index=False, float_format=lambda x: f"{x:,.3f}"))
    best = min(lcohs, key=lambda k: lcohs[k]["LCOH"])
    print(f"\n  >>> 최저 LCOH : {best}  ({lcohs[best]['LCOH']:.3f} USD/kg-H2)")

    print("\n[손실 원장 3케이스 대조]")
    led = pd.DataFrame({name: {k: v / r.E_paid * 100 for k, v in r.ledger.items()}
                        for name, r in results.items()}).T
    led.columns = [f"{bucket} [%]" for bucket in led.columns]
    print(led.to_string(float_format=lambda x: f"{x:,.2f}"))

    sweeps = {}
    if make_plots and plt is not None:
        _banner("그래프 생성 (G01 ~ G20)")
        sweeps = make_all_figures(profile, results, lcohs, fixed_stack,
                                  out1, out2, out3)
        print(f"  figures: {_ensure_fig_dir() or '(저장 안 함)'}")

    summary = print_final_summary(profile, fixed_stack, results, lcohs,
                                  out1, out2, out3, sweeps)

    files = export_multiyear_results(profile, results, lcohs,
        {"Case 1": out1, "Case 2": out2, "Case 3": out3}, cmp_df, sweeps)
    summary_excel = summary.attrs.get("excel_path")
    if summary_excel:
        files[Path(summary_excel).name] = summary_excel
    return {"profile": profile, "stack": fixed_stack, "window": win, "saved_files": files,
            "summary_excel": summary_excel,
            "case1": out1, "case2": out2, "case3": out3, "results": results,
            "lcohs": lcohs, "compare": cmp_df, "degradation": degradation_df, "summary": summary,
            "ess_summary": ess_df, "lifecycle_summary": life_df, "j_sweeps": sweeps,
            "rating_sweep": rating_df,
            "j_headline": j_headline}


def sensitivity_stack_fraction(fractions=(0.25, 0.4, 0.55, 0.7, 1.0),
                               verbose: bool = True) -> pd.DataFrame:
    global STACK_PV_FRACTION
    profile = get_profile()
    backup, rows = STACK_PV_FRACTION, []
    try:
        for f in fractions:
            STACK_PV_FRACTION = f
            st = build_fixed_stack(profile, verbose=False)
            o1 = run_case1(profile, st, verbose=False)
            o2 = run_case2(profile, st, verbose=False)
            o3 = run_case3(profile, st, verbose=False)
            lc = {"C1": o1["lcoh"]["LCOH"], "C2": o2["lcoh"]["LCOH"],
                  "C3": o3["lcoh"]["LCOH"]}
            rows.append({"PV fraction": f, "nameplate [kW]": st.P_nameplate,
                         "C1 curtail [%]": o1["result"].E_curtail / o1["result"].E_paid * 100,
                         "C1 H2 [t]": _annual_mean(o1["result"], o1["result"].H2_total) / 1e3,
                         "C2 H2 [t]": _annual_mean(o2["result"], o2["result"].H2_total) / 1e3,
                         "C3 H2 [t]": _annual_mean(o3["result"], o3["result"].H2_total) / 1e3,
                         "C1 LCOH": lc["C1"], "C2 LCOH": lc["C2"], "C3 LCOH": lc["C3"],
                         "winner": min(lc.items(), key=lambda kv: kv[1])[0]})
            if verbose:
                print(f"  fraction {f:.2f}: nameplate {st.P_nameplate:,.0f} kW done")
    finally:
        STACK_PV_FRACTION = backup
    df = pd.DataFrame(rows)
    if verbose:
        _banner("민감도 — 스택 규모 vs 재생 규모")
        print(df.to_string(index=False, float_format=lambda x: f"{x:,.3f}"))
    return df


def sensitivity_stack_lifetime(lifetimes=(20_000, 40_000, 60_000, 80_000, 100_000),
                               case: str = "Case 3", verbose: bool = True) -> pd.DataFrame:
    global STACK_LIFETIME_H
    profile, fixed_stack, _ = setup_common(verbose=False)
    backup, rows = STACK_LIFETIME_H, []
    runner = {"Case 1": run_case1, "Case 2": run_case2, "Case 3": run_case3}[case]
    try:
        for life_h in lifetimes:
            STACK_LIFETIME_H = float(life_h)
            out = runner(profile, fixed_stack, verbose=False)
            lc = out["lcoh"]
            rows.append({
                "stack lifetime [h]": life_h,
                "앵커 uV/h": deg_anchor_rate_uV_h(),
                "유효 uV/h": lc.get("degradation", {}).get("rate_eff_uV_h", np.nan),
                "실효 수명 [h]": lc.get("stack_life_hours", np.nan),
                "교체 [회]": lc.get("replacement_count", np.nan),
                "생애 H2 [t]": lc.get("H2_total_life_kg", np.nan) / 1e3,
                "평균 유지율 [%]": lc.get("retention_mean_pct", np.nan),
                "교체비 [$/kg]": lc["items_usd_per_kg"].get("stack_replacement", np.nan),
                "LCOH [$/kg]": lc["LCOH"],
            })
            if verbose:
                print(f"  lifetime {life_h:,.0f} h -> LCOH {lc['LCOH']:.3f} $/kg, "
                      f"교체 {lc.get('replacement_count', 0)}회")
    finally:
        STACK_LIFETIME_H = backup
    df = pd.DataFrame(rows)
    if verbose:
        _banner(f"민감도 — 스택 수명 ({case})")
        print(df.to_string(index=False, float_format=lambda x: f"{x:,.3f}"))
    if plt is not None and len(df):
        setup_plots()
        fig, ax = plt.subplots(figsize=(8.0, 4.4))
        ax.plot(df["stack lifetime [h]"] / 1e3, df["LCOH [$/kg]"], "o-",
                color=CASE_COLORS.get(case, "#333333"), lw=2)
        ax.set_xlabel("Stack lifetime [kh]")
        ax.set_ylabel("LCOH [USD/kg-H$_2$]")
        ax.set_title(f"{case} — stack-lifetime sensitivity (degradation + replacement)")
        ax2 = ax.twinx()
        ax2.bar(df["stack lifetime [h]"] / 1e3, df["교체 [회]"], width=6.0,
                color="#9467bd", alpha=0.35, label="replacements")
        ax2.set_ylabel("Stack replacements in project life [-]")
        ax2.grid(False)
        fig.tight_layout()
        _finish(fig, "S1_stack_lifetime_sensitivity")
    return df


if __name__ == "__main__":
    pd.set_option("display.width", 200)
    pd.set_option("display.max_columns", 50)
    _out = main()

 # %%
